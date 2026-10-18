"""Reconciliation and content checks for the v2 dashboard. Run: python tests/test_dashboard.py  (or python -m pytest tests)

The dashboard is recomputed from the source file (pandas) and reconciled to the v1 MIS workbook's LibreOffice-calculated values.
No Microsoft Excel Desktop test is performed.
"""
import json, os, re, subprocess, sys, tempfile
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'data', 'source', 'Nova_CFO_MIS_Complete_Synthetic_Dataset.xlsx'); OUT = os.path.join(ROOT, 'outputs', 'v2')
HTML = os.path.join(OUT, 'Nova_FPA_Dashboard.html'); DATA = os.path.join(OUT, 'dashboard_data.json')
D = json.load(open(DATA)); H = open(HTML, encoding='utf-8').read()
WB = openpyxl.load_workbook(os.path.join(ROOT, 'CFO_Monthly_Sales_MIS.xlsx'), data_only=True)
close = lambda a, b, t=0.01: abs(a - b) <= t

def test_rebuild_is_deterministic_and_committed_output_current():
    with tempfile.TemporaryDirectory() as t:
        subprocess.run([sys.executable, os.path.join(ROOT, 'src', 'build_dashboard.py'), '--input', SRC, '--out', t], check=True, capture_output=True)
        for f in ('Nova_FPA_Dashboard.html', 'dashboard_data.json'): assert open(os.path.join(t, f), 'rb').read() == open(os.path.join(OUT, f), 'rb').read(), f'{f} is stale: rebuild and review'

def test_monthly_matches_v1_workbook():
    ws = WB['Monthly_MIS']; rows = {ws.cell(r, 1).value: (ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value) for r in range(9, 30)}
    assert len(D['monthly']) == 21 == len(rows)
    for m in D['monthly']: assert all(close(a, b) for a, b in zip((m['revenue'], m['cogs'], m['gp']), rows[m['month']])), m['month']
    tot = [ws.cell(30, c).value for c in (5, 6, 7)]; assert all(close(a, b) for a, b in zip((D['total']['revenue'], D['total']['cogs'], D['total']['gp']), tot))
    assert close(sum(m['revenue'] for m in D['monthly']), D['total']['revenue'])

def test_ytd_matches_v1_workbook():
    ws = WB['Monthly_MIS']; lab = {ws.cell(r, 1).value: r for r in range(30, 40) if ws.cell(r, 1).value}
    cur, pri = lab['YTD 2026 (Jan-Sep)'], lab['YTD 2025 (Jan-Sep)']; assert D['current_year'] == 2026 and D['ytd_through_month'] == 9
    for key, r in (('ytd_current', cur), ('ytd_prior', pri)): assert all(close(D[key][k], ws.cell(r, c).value) for k, c in (('revenue', 5), ('cogs', 6), ('gp', 7))), key
    assert close(D['ytd_change']['revenue'], ws.cell(lab['YTD change (current vs prior)'], 5).value, 1e-9)

def test_regions_and_products_match_v1_workbook_and_sum_to_total():
    for sheet, key, name_col, off in (('Region_Performance', 'regions', 1, 0), ('Product_Performance', 'products', 1, 1)):
        ws = WB[sheet]; v1 = {}
        for r in range(1, 40):
            if ws.cell(r, 1).value in ('Region', 'Product'): first = r + 1; break
        for r in range(first, first + 12):
            n = ws.cell(r, 1).value
            if n == 'Total' or n is None: break
            v1[n] = [ws.cell(r, c).value for c in (2 + off, 3 + off, 4 + off, 5 + off, 7 + off, 8 + off)]
        assert len(v1) == len(D[key]), sheet
        for row in D[key]:
            e = v1[row['name']]
            assert close(row['revenue'], e[0]) and close(row['cogs'], e[1]) and close(row['gp'], e[2]) and close(row['gm'], e[3], 1e-9) and close(row['ytd_revenue'], e[4]) and close(row['ytd_revenue_prior'], e[5]), (sheet, row['name'])
        assert close(sum(r['revenue'] for r in D[key]), D['total']['revenue']), key

def test_overlap_facts_and_budget_diagnostics_match_v1_fact_pack():
    assert D['overlap']['max_regions_per_product'] == 1 and D['overlap']['max_products_per_region'] == 2
    facts = ' '.join(json.load(open(os.path.join(ROOT, 'fact_pack.json')))['facts'])
    assert f"{D['budget']['ratio_vs_actual']:.1f}x" in facts and f"{D['budget']['share_in_cells_without_sales'] * 100:.0f}%" in facts
    assert D['source_sha256'] == open(os.path.join(ROOT, 'data', 'source', 'SHA256SUMS')).read().split()[0]

def test_required_disclosures_present():
    for must in ('DATA-QUALITY WARNING', 'budget is NOT comparable', 'shows no budget-vs-actual variance', 'illustrative only', 'not an approved budget', 'region and product effects overlap',
                 'confounded', 'no live refresh, no autonomous publishing', 'Not validated in Excel Desktop', 'Not production-ready', 'separate from Provisional v1', 'Synthetic data', 'How to refresh'):
        assert must.lower() in H.lower(), must
    assert H.index('DATA-QUALITY WARNING') < H.index('<div class="kpis">'), 'warning must appear before the KPIs'

def test_no_budget_variance_numbers_or_unsupported_claims():
    assert 'Budget_v2' in H and 'variance' in H.lower()
    for bad in ('live data', 'auto-refresh', 'refreshes automatically', 'validated in excel desktop', 'production-ready dashboard', 'approved release'):
        assert bad not in H.lower().replace('not validated in excel desktop', '').replace('not production-ready', ''), bad
    assert 'Not validated in Excel Desktop' in H

def test_every_bar_carries_a_value_label():
    import re as _re
    bars = _re.findall(r'<rect [^>]*class="(bar|bar-prior|seg-cogs|seg-gp)"', H)
    labels = len(_re.findall(r'<text [^>]*class="(?:vl|vli dk|vli)"', H))
    # 42 stacked segments (+21 revenue totals above them = 63 labels), 6 YTD bars (6 labels), 12 region/product columns (value above + margin inside = 24); plus 21 margin-line point labels below
    assert len(bars) == 42 + 6 + 4 + 8 and labels == 21 * 3 + 6 + 2 * (4 + 8), (len(bars), labels)
    assert H.count('class="vl sm"') == 21

def test_self_contained_static_file():
    assert '<script' not in H.lower() and 'src=' not in H.lower() and not re.search(r'https?://', H) and '@import' not in H and 'url(' not in H

if __name__ == '__main__':
    for n, f in list(globals().items()):
        if n.startswith('test_'): f(); print('PASS', n)
