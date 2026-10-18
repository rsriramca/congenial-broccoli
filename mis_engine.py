#!/usr/bin/env python3
"""Nova Consumer Products - CFO Monthly Sales MIS engine (v2: budget integrity + synthetic v2 budget).

Usage:  python build_cfo_mis.py --input <dataset.xlsx> --output CFO_Monthly_Sales_MIS.xlsx [--validation validation.json]
Reads the source workbook (never modifies it), copies the source tables into the output unchanged,
and builds formula-driven reports. Actual = rows with Status == 'Actual' (cross-checked to Calendar.Reporting_Status).
"""
import argparse, json, re, sys
import pandas as pd, numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule

REQ = {
 'Sales_Transactions': ['Invoice_ID','Month','Status','Customer_ID','Product_ID','Units','Net_Price_INR','Unit_Cost_INR','Net_Revenue_INR','COGS_INR','Gross_Profit_INR'],
 'Product_Master': ['Product_ID','Product_Name','Category','List_Price_INR','Standard_Cost_INR'],
 'Customer_Master': ['Customer_ID','Customer_Name','Region','Channel'],
 'Calendar': ['Month','Year','Quarter','Month_Name','Reporting_Status'],
}
OPT = ['Sales_Budget']
INR = r'[>=10000000]##\,##\,##\,##0;[>=100000]##\,##\,##0;##,##0'
PCT = '0.0%'; PP = '+0.0" pp";-0.0" pp";0.0" pp"'; PCTS = '+0.0%;-0.0%;0.0%'
NAVY='1F3864'; TEAL='2E8B8B'; ORANGE='C55A11'; GREY='F2F2F2'
F = 'Arial'
fh = Font(name=F, bold=True, color='FFFFFF', size=10); fb = Font(name=F, size=10); fbb = Font(name=F, size=10, bold=True)
ft = Font(name=F, size=16, bold=True, color=NAVY); fs = Font(name=F, size=10, italic=True, color='595959')
fin = Font(name=F, size=10, color='0000FF')  # hardcoded inputs / baselines
hfill = PatternFill('solid', fgColor=NAVY); gfill = PatternFill('solid', fgColor=GREY); tfill = PatternFill('solid', fgColor='D9E2F3')
thin = Side(style='thin', color='BFBFBF'); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical='top')

class EngineError(Exception): pass

def die(msg): raise EngineError(msg)

def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col+i, t); c.font = fh; c.fill = hfill; c.border = box
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 32

def put(ws, row, col, v, fmt=None, font=None, fill=None, al=None, border=True):
    c = ws.cell(row, col, v); c.font = font or fb
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if border: c.border = box
    return c

def title(ws, t, sub):
    ws.sheet_view.showGridLines = False
    ws['A1'] = t; ws['A1'].font = ft
    ws['A2'] = sub; ws['A2'].font = fs

def widths(ws, ws_w):
    for k, v in ws_w.items(): ws.column_dimensions[k].width = v

# ------------------------------------------------------------------ load + schema validation
def load(path):
    x = pd.read_excel(path, sheet_name=None)
    for sh, cols in REQ.items():
        if sh not in x: die(f'Missing required sheet {sh}')
        miss = [c for c in cols if c not in x[sh].columns]
        if miss: die(f'Sheet {sh} is missing columns {miss}')
    s = x['Sales_Transactions']
    bad = set(s.Status.dropna().unique()) - {'Actual','Forecast'}
    if bad: die(f'Unexpected Status values in Sales_Transactions: {bad}')
    return x

def source_baselines(x):
    b = {}
    cc = x.get('Control_Checks')
    if cc is not None and {'Control','Expected_Value'} <= set(cc.columns):
        b = dict(zip(cc.Control, cc.Expected_Value))
    return b

# ------------------------------------------------------------------ independent pandas computation (used for commentary + exception snapshot)
def pandas_facts(x):
    s = x['Sales_Transactions'].merge(x['Customer_Master'], on='Customer_ID', how='left').merge(x['Product_Master'], on='Product_ID', how='left')
    a = s[s.Status=='Actual'].copy()
    a['y'] = a.Month.str[:4].astype(int); a['m'] = a.Month.str[5:].astype(int)
    ly = a.y.max(); lm = a[a.y==ly].m.max()
    cur = a[(a.y==ly)&(a.m<=lm)]; pri = a[(a.y==ly-1)&(a.m<=lm)]
    def t(d): r, g = d.Net_Revenue_INR.sum(), d.Gross_Profit_INR.sum(); return r, g, g/r if r else np.nan
    mm = a.groupby('Month')[['Net_Revenue_INR','Gross_Profit_INR']].sum(); mm['gm'] = mm.Gross_Profit_INR/mm.Net_Revenue_INR
    r2 = lambda v: float(np.corrcoef(np.arange(len(v)), v)[0,1]**2)
    def grp(k):
        g = a.groupby(k)[['Net_Revenue_INR','Gross_Profit_INR']].sum(); g['gm'] = g.Gross_Profit_INR/g.Net_Revenue_INR
        gc = cur.groupby(k).Net_Revenue_INR.sum(); gp = pri.groupby(k).Net_Revenue_INR.sum(); g['yoy'] = gc/gp-1
        return g
    return dict(a=a, ly=int(ly), lm=int(lm), tot=t(a), cur=t(cur), pri=t(pri), mm=mm,
                r2rev=r2(mm.Net_Revenue_INR.values), r2gm=r2(mm.gm.values), reg=grp('Region'), prod=grp('Product_Name'),
                latest=f'{ly}-{lm:02d}', first=a.Month.min())

# ------------------------------------------------------------------ facts, budget diagnostics, commentary (v2)
def extra_facts(x, fx):
    a = fx['a']
    ad = a.dropna(subset=['Region', 'Product_Name'])   # rows with unmatched keys are reported by the controls, not here
    rp = {k: sorted(v) for k, v in ad.groupby('Region').Product_Name.unique().items()}
    pr = ad.groupby('Product_Name').Region.nunique().to_dict()
    return dict(reg_products=rp, prod_nregions=pr, confounded=bool(pr) and max(pr.values()) == 1)

def budget_diagnostics(x, fx):
    if 'Sales_Budget' not in x: return None
    b = x['Sales_Budget'].copy(); cal = x['Calendar']; a = fx['a']; mm = fx['mm']; tr = fx['tot'][0]
    d = {}
    d['months_aligned'] = set(b.Month) == set(cal.Month)
    cells = sorted(set(zip(b.Region, b.Category))); cov = set(zip(a.Region, a.Category)); unc = [c for c in cells if c not in cov]
    am = b[b.Month.isin(mm.index)].copy(); am['incov'] = [(r, c) in cov for r, c in zip(am.Region, am.Category)]
    bud_all = am.Revenue_Budget_INR.sum(); bud_in = am[am.incov].Revenue_Budget_INR.sum()
    d.update(cells=cells, uncovered=unc, n_cells=len(cells), n_cov=len(cov), ratio_all=bud_all/tr, ratio_in=bud_in/tr, unc_share=1-bud_in/bud_all)
    mc = am[am.incov].merge(a.groupby(['Month', 'Region', 'Category']).Net_Revenue_INR.sum().reset_index(), on=['Month', 'Region', 'Category'], how='left')
    r = (mc.Revenue_Budget_INR/mc.Net_Revenue_INR).dropna()
    d.update(cm_min=float(r.min()), cm_max=float(r.max()), cm_cv=float(r.std()/r.mean()), cm_med=float(r.median()), cm_n=int(len(r)))
    d['ann_ratio'] = bud_all/12/tr; d['qtr_ratio'] = bud_all/3/tr
    d['avg_inv'] = float(a.Net_Revenue_INR.mean()); d['inv_per_month'] = float(len(a)/len(mm)); d['budget_per_month'] = bud_all/len(mm)
    d['inv_needed'] = d['budget_per_month']/d['avg_inv']; d['cust_inv'] = int(a.Customer_ID.nunique()); d['cust_all'] = int(len(x['Customer_Master']))
    # generator signature
    const_in_year = all(g.Revenue_Budget_INR.nunique() == 1 for _, g in b.assign(y=b.Month.str[:4]).groupby(['y', 'Region', 'Category']))
    gp_pcts = (b.Gross_Profit_Budget_INR/b.Revenue_Budget_INR).round(6).unique()
    yrs = sorted(b.Month.str[:4].unique()); up = None
    if len(yrs) > 1:
        m0_ = b[b.Month == b.Month[b.Month.str[:4] == yrs[0]].min()].set_index(['Region', 'Category']).Revenue_Budget_INR
        m1_ = b[b.Month == b.Month[b.Month.str[:4] == yrs[1]].min()].set_index(['Region', 'Category']).Revenue_Budget_INR
        u = (m1_/m0_).round(6).unique(); up = float(u[0]) if len(u) == 1 else None
    g0 = b[b.Month == b.Month.min()].pivot_table(index='Region', columns='Category', values='Revenue_Budget_INR', aggfunc='first')
    resid = float((g0 - (g0.mean(axis=1).values[:, None] + g0.mean(axis=0).values[None, :] - g0.values.mean())).abs().values.max())
    d.update(const_in_year=bool(const_in_year), gp_pcts=[float(v) for v in gp_pcts], uplift=up, add_resid=resid,
             grid_exact=bool(const_in_year and len(gp_pcts) == 1 and resid < 1))
    d['bud_gp'] = float(b.Gross_Profit_Budget_INR.sum()/b.Revenue_Budget_INR.sum()); d['act_gp'] = fx['tot'][2]
    # calibration year = latest actual year - 1, only if fully actual
    calact = x['Calendar'][x['Calendar'].Reporting_Status == 'Actual'].Month
    cy = fx['ly']-1; full = sum(1 for m in calact if m.startswith(str(cy))) == 12
    d['calib_year'] = cy if full else None
    d['flag'] = bool(d['ratio_all'] > 1.5 or d['ratio_all'] < 0.67 or len(unc) > 0)
    d['v2'] = bool(d['flag'] and d['calib_year'])
    if d['v2']:
        bb = b[b.Month.str[:4] == str(cy)]; bb = bb[[(r, c) in cov for r, c in zip(bb.Region, bb.Category)]]
        d['k'] = float(a[a.Month.str[:4] == str(cy)].Net_Revenue_INR.sum()/bb.Revenue_Budget_INR.sum())
    f = lambda v: f'{v:,.1f}'
    T = []
    T.append(('Reporting periods', 'Set of Sales_Budget months vs Calendar months', 'Aligned: same 24 months' if d['months_aligned'] else 'MISMATCH in month sets', 'NO PERIOD MISMATCH' if d['months_aligned'] else 'CONFIRMED ISSUE'))
    T.append(('Aggregation level', 'Budget rows per month; grain of each table', f"Budget = Month x Region x Category ({d['n_cells']} cells/month, no product or customer dimension). Transactions = invoice x product, aggregated here to the same grain. Product-level budget cannot be derived.", 'COMPARABLE ONLY AT REGION x CATEGORY'))
    T.append(('Units / currency / scale factor', 'Budget/actual ratio for each in-scope Month x Region x Category cell', f"{d['cm_n']} cell-months: ratio ranges {d['cm_min']:.1f}x to {d['cm_max']:.1f}x (median {d['cm_med']:.1f}x, CV {d['cm_cv']:.0%}). A uniform unit or currency conversion would give one constant ratio.", 'RULED OUT as sole cause'))
    T.append(('Periodicity (annual / quarterly figures)', 'Budget divided by 12 or by 3 vs actual', f"Budget/12 = {d['ann_ratio']:.2f}x actual; Budget/3 = {d['qtr_ratio']:.2f}x. Neither is close to 1.0x.", 'RULED OUT'))
    T.append(('Business scope (coverage of Region x Category cells)', 'Cells with budget but no actual sales in any month', f"{len(unc)} of {d['n_cells']} cells have zero sales in the source ({', '.join(f'{r}/{c}' for r, c in unc)}); they hold {d['unc_share']:.1%} of budget for the actual months. Cause in the data: each product is sold in one region only, so each region covers only {len(cov)}/{d['n_cells']} of the grid. Restricting budget to cells that have sales reduces the gap from {d['ratio_all']:.2f}x to {d['ratio_in']:.2f}x.", 'CONFIRMED (explains part of the gap)'))
    T.append(('Transaction coverage', 'Invoices per month, customers invoiced, avg invoice value', f"{d['inv_per_month']:.0f} invoices/month, avg INR {d['avg_inv']:,.0f}; {d['cust_inv']}/{d['cust_all']} customers invoiced. Budget of INR {d['budget_per_month']:,.0f}/month implies about {d['inv_needed']:.0f} invoices/month at that average. Cannot tell whether the transaction table is a sample or the budget is overstated.", 'UNRESOLVED'))
    T.append(('Data-generation signature', 'Fit of budget to a closed-form grid', (f"Budget is constant within each year (no seasonality), gross margin is exactly {d['gp_pcts'][0]:.0%} in every row, 2nd-year uplift is exactly {up:.2f}x for every cell, and revenue is additive in Region + Category effects (max residual INR {resid:.2f}). It is formula-generated, not built from transactions." if d['grid_exact'] and up else 'Budget does not fit a simple closed-form grid.'), 'CONFIRMED' if d['grid_exact'] else 'NOT CONFIRMED'))
    T.append(('Margin and trend consistency', 'Budget GP % vs actual GP %; budget growth vs actual growth', f"Budget GP % {d['bud_gp']:.1%} vs actual {d['act_gp']:.1%}. Budget steps up {((up-1) if up else 0):.0%} in year 2; actual YTD revenue moved {fx['cur'][0]/fx['pri'][0]-1:+.1%}.", 'INFORMATION'))
    d['tests'] = T
    d['root'] = (f"PARTIALLY DEMONSTRATED. Demonstrated: (1) the budget is a closed-form formula grid with no link to the transaction table; (2) {d['unc_share']:.0%} of it sits in {len(unc)} Region x Category cells where the source has no sales at all, because each product is sold in one region only. "
                 f"Removing those cells narrows the gap from {d['ratio_all']:.1f}x to {d['ratio_in']:.1f}x. NOT demonstrated: the remaining {d['ratio_in']:.1f}x. It is not a constant factor ({d['cm_min']:.1f}x-{d['cm_max']:.1f}x), so it is not a unit/currency issue. "
                 "Most likely the two tables were generated independently without calibration, but that cannot be proven from the data. Status: UNRESOLVED - business confirmation needed.")
    return d

def _fm(v, kind):
    return {'cr': f'{v/1e7:.2f} crore', 'lk': f'{v/1e5:.1f}', 'p1': f'{v:.1%}', 'sp': f'{v:+.1%}'}[kind]

def fact_pack(x, fx, bd, ef):
    lines = []; tot = fx['tot']; cur = fx['cur']; pri = fx['pri']; mm = fx['mm']
    lines += [f"Total actual net revenue INR {_fm(tot[0],'cr')} ({tot[0]/1e5:.1f} lakh)", f"Total gross profit INR {_fm(tot[1],'cr')}", f"Total margin {_fm(tot[2],'p1')}",
              f"YTD revenue INR {_fm(cur[0],'cr')}; prior-year YTD INR {_fm(pri[0],'cr')}; change {_fm(cur[0]/pri[0]-1,'sp')}", f"YTD margin {_fm(cur[2],'p1')} vs {_fm(pri[2],'p1')}",
              f"Monthly revenue range INR {_fm(mm.Net_Revenue_INR.min(),'lk')}-{_fm(mm.Net_Revenue_INR.max(),'lk')} lakh", f"Monthly margin range {_fm(mm.gm.min(),'p1')}-{_fm(mm.gm.max(),'p1')}",
              f"{len(mm)} months; {len(fx['reg'])} regions; {len(fx['prod'])} products"]
    for k, r in fx['reg'].iterrows(): lines.append(f"Region {k}: margin {_fm(r.gm,'p1')}; YTD revenue change {_fm(r.yoy,'sp')}; revenue INR {_fm(r.Net_Revenue_INR,'lk')} lakh")
    for k, r in fx['prod'].iterrows(): lines.append(f"Product {k}: margin {_fm(r.gm,'p1')}; YTD revenue change {_fm(r.yoy,'sp')}; revenue INR {_fm(r.Net_Revenue_INR,'lk')} lakh")
    if bd:
        lines += [f"Budget/actual ratio {bd['ratio_all']:.0f}x ({bd['ratio_all']:.1f}x, {bd['ratio_all']:.2f}x); in-scope {bd['ratio_in']:.0f}x ({bd['ratio_in']:.1f}x, {bd['ratio_in']:.2f}x)", f"Budget in cells with no sales {bd['unc_share']:.0%} ({bd['unc_share']:.1%}); {len(bd['uncovered'])} of {bd['n_cells']} cells"]
    toks = set()
    for l in lines: toks |= set(_numtok(l))
    toks |= {str(y) for y in range(2020, 2036)}
    return dict(facts=lines, allowed=sorted(toks))

def _numtok(t):
    return [m.replace(',', '').rstrip('.') for m in re.findall(r'\d[\d,]*(?:\.\d+)?%?', t)]

def commentary(x, fx, bd, ef):
    tr, tg, tm = fx['tot']; cr_, cg, cm = fx['cur']; pr_, pg, pm = fx['pri']
    mm = fx['mm']; reg, prod = fx['reg'], fx['prod']
    cr = lambda v: f'{v/1e7:.2f} crore'; lk = lambda v: f'{v/1e5:.1f}'
    trend = 'no clear trend' if max(fx['r2rev'], fx['r2gm']) < 0.1 else 'a trend worth review (see Monthly_MIS)'
    mon = pd.Timestamp(fx['latest']+'-01').strftime('%b-%Y'); first = pd.Timestamp(fx['first']+'-01').strftime('%b-%Y'); n = len(mm)
    fall_r = reg.yoy.idxmin(); worst_r = reg.gm.idxmin()
    obs = (f"OBSERVED (actuals {first} to {mon}, {n} months): Net revenue INR {cr(tr)}, gross profit INR {cr(tg)}, margin {tm:.1%}. "
           f"YTD {mon} revenue INR {cr(cr_)}, {cr_/pr_-1:+.1%} vs prior-year YTD; margin {cm:.1%} vs {pm:.1%}. "
           f"Monthly revenue (INR {lk(mm.Net_Revenue_INR.min())}-{lk(mm.Net_Revenue_INR.max())} lakh) and margin ({mm.gm.min():.1%}-{mm.gm.max():.1%}) show {trend}. ")
    if ef['confounded']:
        ps = sorted(ef['reg_products'][fall_r], key=lambda p: prod.yoy[p])
        pw = sorted(ef['reg_products'][worst_r], key=lambda p: prod.gm[p])
        obs += (f"Each product sells in one region only, so region and product effects overlap: {fall_r}'s {reg.yoy[fall_r]:+.1%} YTD is "
                + ' against '.join(f"{p} ({prod.yoy[p]:+.1%})" for p in ps) + f"; {worst_r}'s lowest margin ({reg.gm[worst_r]:.1%}) reflects "
                + ' and '.join(f"{p} ({prod.gm[p]:.1%})" for p in pw) + '.')
    else:
        fp = prod.yoy.idxmin()
        obs += (f"{fall_r} revenue moved {reg.yoy[fall_r]:+.1%} YTD; {fp} {prod.yoy[fp]:+.1%}. {worst_r} has the lowest margin ({reg.gm[worst_r]:.1%}).")
    unv = "UNVERIFIED CAUSES: Drivers of these movements (pricing, promotions, distribution, input costs, mix) cannot be established from this data. Returns, rebates and GST are excluded."
    if bd and bd['flag']:
        unv += f" The original budget is not comparable (about {bd['ratio_all']:.0f}x actuals; {bd['unc_share']:.0%} sits in cells with no sales), so no approved budget variance is reported."
    return obs, unv

def check_commentary(text, pack):
    probs = []; w = len(text.split())
    if w > 150: probs.append(f'{w} words (limit 150)')
    if 'OBSERVED' not in text or 'UNVERIFIED' not in text: probs.append('must contain separate OBSERVED and UNVERIFIED CAUSES sections')
    allowed = set(pack['allowed']); bad = [t for t in _numtok(text) if t not in allowed]
    if bad: probs.append(f'numbers not found in verified fact pack: {sorted(set(bad))}')
    return probs

# ------------------------------------------------------------------ workbook build
def build(inp, out, validation=None, commentary_override=None):
    x = load(inp); base = source_baselines(x); fx = pandas_facts(x); ef = extra_facts(x, fx); bd = budget_diagnostics(x, fx)
    wb = Workbook(); wb.remove(wb.active)
    sheets = ['Executive_Summary','Monthly_MIS','Region_Performance','Product_Performance'] + (['Budget_vs_Actual_v2'] if (bd and bd['v2']) else []) + (['Budget_Integrity'] if bd else []) + (['Budget_v2_Synthetic'] if (bd and bd['v2']) else []) + ['Control_Checks','Validation_Report',
              'Data_Dictionary','Automation_Readme','Calc_Actuals','Sales_Transactions','Product_Master','Customer_Master','Calendar']
    if 'Sales_Budget' in x: sheets.append('Sales_Budget')
    W = {n: wb.create_sheet(n) for n in sheets}

    # ---- raw copies (values exactly as in source, original sheet names)
    raw_dim = {}
    for sh in [k for k in REQ] + [k for k in OPT if k in x]:
        ws = W[sh]; df = x[sh]
        ws.append(list(df.columns))
        for rec in df.itertuples(index=False): ws.append([None if (isinstance(v,float) and np.isnan(v)) else (v.item() if hasattr(v,'item') else v) for v in rec])
        for c in ws[1]: c.font = fh; c.fill = hfill
        for col in range(1, len(df.columns)+1): ws.column_dimensions[L(col)].width = 18
        ws.freeze_panes = 'A2'; ws.sheet_properties.tabColor = '7F7F7F'
        raw_dim[sh] = len(df)+1
    nS = raw_dim['Sales_Transactions']; nP = raw_dim['Product_Master']; nC = raw_dim['Customer_Master']; nK = raw_dim['Calendar']
    ST = lambda c: f"Sales_Transactions!${c}$2:${c}${nS}"
    PM = lambda c: f"Product_Master!${c}$2:${c}${nP}"
    CM = lambda c: f"Customer_Master!${c}$2:${c}${nC}"
    CAL = lambda c: f"Calendar!${c}$2:${c}${nK}"
    CA = lambda c: f"Calc_Actuals!${c}$2:${c}${nS}"

    # ---- Calc_Actuals
    ws = W['Calc_Actuals']; ws.sheet_properties.tabColor = '7F7F7F'
    cols = ['Invoice_ID','Month','Status','Customer_ID','Product_ID','Region','Channel','Category','Product_Name','Year','Month_No','Calendar_Status',
            'Units','Net_Revenue_INR','COGS_INR','Gross_Profit_INR','Flag_Unmatched_Product','Flag_Unmatched_Customer','Flag_Unmatched_Month',
            'Flag_Duplicate_Invoice','Flag_Arithmetic','Flag_Status_vs_Calendar','Any_Exception']
    hdr(ws, 1, cols)
    for r in range(2, nS+1):
        f = {
         1:f"=Sales_Transactions!A{r}", 2:f"=Sales_Transactions!B{r}", 3:f"=Sales_Transactions!C{r}", 4:f"=Sales_Transactions!D{r}", 5:f"=Sales_Transactions!E{r}",
         6:f'=IFERROR(INDEX({CM("C")},MATCH(D{r},{CM("A")},0)),"UNMATCHED")', 7:f'=IFERROR(INDEX({CM("D")},MATCH(D{r},{CM("A")},0)),"UNMATCHED")',
         8:f'=IFERROR(INDEX({PM("C")},MATCH(E{r},{PM("A")},0)),"UNMATCHED")', 9:f'=IFERROR(INDEX({PM("B")},MATCH(E{r},{PM("A")},0)),"UNMATCHED")',
         10:f'=IFERROR(INDEX({CAL("B")},MATCH(B{r},{CAL("A")},0)),0)', 11:f'=IFERROR(VALUE(RIGHT(B{r},2)),0)',
         12:f'=IFERROR(INDEX({CAL("E")},MATCH(B{r},{CAL("A")},0)),"UNMATCHED")',
         13:f"=Sales_Transactions!F{r}", 14:f"=Sales_Transactions!I{r}", 15:f"=Sales_Transactions!J{r}", 16:f"=Sales_Transactions!K{r}",
         17:f'=IF(ISNUMBER(MATCH(E{r},{PM("A")},0)),0,1)', 18:f'=IF(ISNUMBER(MATCH(D{r},{CM("A")},0)),0,1)', 19:f'=IF(ISNUMBER(MATCH(B{r},{CAL("A")},0)),0,1)',
         20:f'=IF(COUNTIF({ST("A")},A{r})>1,1,0)',
         21:f'=IF(OR(ABS(Sales_Transactions!F{r}*Sales_Transactions!G{r}-Sales_Transactions!I{r})>0.01,ABS(Sales_Transactions!F{r}*Sales_Transactions!H{r}-Sales_Transactions!J{r})>0.01,ABS(Sales_Transactions!I{r}-Sales_Transactions!J{r}-Sales_Transactions!K{r})>0.01),1,0)',
         22:f'=IF(C{r}<>L{r},1,0)', 23:f'=IF(SUM(Q{r}:V{r})>0,1,0)'}
        for c, v in f.items():
            cell = ws.cell(r, c, v); cell.font = fb
            if c in (14,15,16): cell.number_format = '#,##0.00'
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f"A1:W{nS}"
    for c in range(1, 24): ws.column_dimensions[L(c)].width = 16

    # ---- Monthly_MIS
    mw = W['Monthly_MIS']; title(mw, 'Monthly Sales & Profitability MIS - Actuals only', 'Nova Consumer Products (synthetic). INR. Net revenue excludes GST, returns, rebates and intercompany (per source README). Forecast and budget are excluded.')
    for i, (lab, fm) in enumerate([('Latest actual month (from Calendar)', f'=LOOKUP(2,1/({CAL("E")}="Actual"),{CAL("A")})'),
                                   ('Latest actual year','=VALUE(LEFT(B4,4))'), ('Latest actual month no.','=VALUE(RIGHT(B4,2))')]):
        put(mw, 4+i, 1, lab, font=fbb, fill=gfill); put(mw, 4+i, 2, fm, font=fb)
    months = list(fx['mm'].index); m0 = 9; m1 = m0+len(months)-1; mt = m1+1
    hdr(mw, 8, ['Month','Year','Invoices','Units','Net Revenue (INR)','COGS (INR)','Gross Profit (INR)','GP %','MoM Revenue %','YoY Revenue %','YoY GP % change (pp)','Check: raw-source revenue','Check: raw-source COGS'])
    rowof = {mo: m0+i for i, mo in enumerate(months)}
    for mo, r in rowof.items():
        y, mn = int(mo[:4]), int(mo[5:])
        put(mw, r, 1, mo); put(mw, r, 2, f'=INDEX({CAL("B")},MATCH(A{r},{CAL("A")},0))')
        MK = f'{CA("J")},B{r},{CA("K")},VALUE(RIGHT(A{r},2))'
        put(mw, r, 3, f'=COUNTIFS({MK},{CA("C")},"Actual")', '#,##0')
        put(mw, r, 4, f'=SUMIFS({CA("M")},{MK},{CA("C")},"Actual")', '#,##0')
        put(mw, r, 5, f'=SUMIFS({CA("N")},{MK},{CA("C")},"Actual")', INR)
        put(mw, r, 6, f'=SUMIFS({CA("O")},{MK},{CA("C")},"Actual")', INR)
        put(mw, r, 7, f'=E{r}-F{r}', INR); put(mw, r, 8, f'=IF(E{r}=0,"n/a",G{r}/E{r})', PCT)
        pm = f'{y}-{mn-1:02d}' if mn > 1 else f'{y-1}-12'; py = f'{y-1}-{mn:02d}'
        put(mw, r, 9, f'=E{r}/E{rowof[pm]}-1' if pm in rowof else 'n/a', PCTS, al=Alignment(horizontal='right'))
        put(mw, r, 10, f'=E{r}/E{rowof[py]}-1' if py in rowof else 'n/a', PCTS, al=Alignment(horizontal='right'))
        put(mw, r, 11, f'=(H{r}-H{rowof[py]})*100' if py in rowof else 'n/a', PP, al=Alignment(horizontal='right'))
        put(mw, r, 12, f'=SUMPRODUCT(({ST("B")}=A{r})*({ST("C")}="Actual")*{ST("I")})', INR, font=fs)
        put(mw, r, 13, f'=SUMPRODUCT(({ST("B")}=A{r})*({ST("C")}="Actual")*{ST("J")})', INR, font=fs)
    put(mw, mt, 1, 'Total (actuals)', font=fbb, fill=tfill)
    for c in (2,): put(mw, mt, c, None, fill=tfill)
    for c, fmt in ((3,'#,##0'),(4,'#,##0'),(5,INR),(6,INR),(7,INR)): put(mw, mt, c, f'=SUM({L(c)}{m0}:{L(c)}{m1})', fmt, font=fbb, fill=tfill)
    put(mw, mt, 8, f'=G{mt}/E{mt}', PCT, font=fbb, fill=tfill)
    for c in (9,10,11): put(mw, mt, c, None, fill=tfill)
    for c in (12,13): put(mw, mt, c, f'=SUM({L(c)}{m0}:{L(c)}{m1})', INR, font=fbb, fill=tfill)
    # period summary
    ps = mt+3; mw.cell(ps-1, 1, 'Like-for-like period summary (actuals)').font = fbb
    hdr(mw, ps, ['Period','','','','Net Revenue (INR)','COGS (INR)','Gross Profit (INR)','GP %'])
    def sumif(col, yr, ytd):
        extra = f',{CA("K")},"<="&$B$6' if ytd else ''
        return f'SUMIFS({CA(col)},{CA("C")},"Actual",{CA("J")},{yr}{extra})'
    per = [('="YTD "&B5&" (Jan-"&TEXT(DATE(B5,B6,1),"mmm")&")"','$B$5',True), ('="YTD "&(B5-1)&" (Jan-"&TEXT(DATE(B5,B6,1),"mmm")&")"','($B$5-1)',True), ('="Full year "&(B5-1)','($B$5-1)',False)]
    for i, (lab, yr, ytd) in enumerate(per):
        r = ps+1+i; put(mw, r, 1, lab, font=fbb)
        for c in (2,3,4): put(mw, r, c, None)
        put(mw, r, 5, '='+sumif('N', yr, ytd), INR); put(mw, r, 6, '='+sumif('O', yr, ytd), INR)
        put(mw, r, 7, f'=E{r}-F{r}', INR); put(mw, r, 8, f'=IF(E{r}=0,"n/a",G{r}/E{r})', PCT)
    r = ps+4; put(mw, r, 1, 'YTD change (current vs prior)', font=fbb, fill=tfill)
    for c in (2,3,4): put(mw, r, c, None, fill=tfill)
    put(mw, r, 5, f'=IF(E{ps+2}=0,"n/a",E{ps+1}/E{ps+2}-1)', PCTS, fill=tfill, font=fbb); put(mw, r, 6, f'=IF(F{ps+2}=0,"n/a",F{ps+1}/F{ps+2}-1)', PCTS, fill=tfill, font=fbb)
    put(mw, r, 7, f'=IF(G{ps+2}=0,"n/a",G{ps+1}/G{ps+2}-1)', PCTS, fill=tfill, font=fbb); put(mw, r, 8, f'=(H{ps+1}-H{ps+2})*100', PP, fill=tfill, font=fbb)
    PS = ps
    widths(mw, {'A':36,'B':11,'C':10,'D':11,'E':18,'F':16,'G':18,'H':9,'I':14,'J':14,'K':16,'L':20,'M':20}); mw.freeze_panes = 'A9'
    # chart
    ch = BarChart(); ch.type = 'col'; ch.title = 'Monthly Net Revenue & Gross Profit (INR) with GP %'; ch.height = 9; ch.width = 26
    ch.add_data(Reference(mw, min_col=5, min_row=8, max_row=m1), titles_from_data=True)
    ch.add_data(Reference(mw, min_col=7, min_row=8, max_row=m1), titles_from_data=True)
    ch.set_categories(Reference(mw, min_col=1, min_row=m0, max_row=m1))
    ch.series[0].graphicalProperties.solidFill = NAVY; ch.series[1].graphicalProperties.solidFill = TEAL
    ch.y_axis.title = 'INR'; ch.y_axis.number_format = '#,##0'; ch.y_axis.majorGridlines = None
    ln = LineChart(); ln.add_data(Reference(mw, min_col=8, min_row=8, max_row=m1), titles_from_data=True)
    ln.y_axis.axId = 200; ln.y_axis.title = 'GP %'; ln.y_axis.number_format = '0%'; ln.y_axis.crosses = 'max'
    ln.y_axis.scaling.min = 0.30; ln.y_axis.scaling.max = 0.40
    ln.series[0].graphicalProperties.line.solidFill = ORANGE; ln.series[0].graphicalProperties.line.width = 28000; ln.series[0].smooth = False
    ch.x_axis.delete = False; ch.y_axis.delete = False; ln.y_axis.delete = False; ch.x_axis.tickLblPos = 'low'
    ch += ln; ch.legend.position = 'b'
    mw.add_chart(ch, f'A{PS+7}')

    # ---- dimension tables (Region / Product)
    def dim_sheet(name, ttl, key_col, keys, keylabel, extra=None):
        ws = W[name]; title(ws, ttl, 'Actuals only. INR. Totals reconcile to Monthly_MIS (see Control_Checks). YTD = Jan to latest actual month; prior-year YTD is like-for-like.')
        heads = [keylabel] + ([extra[0]] if extra else []) + ['Net Revenue (INR)','COGS (INR)','Gross Profit (INR)','GP %','Revenue share','YTD Revenue (current)','YTD Revenue (prior)','YTD Rev YoY %','YTD GP % (current)','YTD GP % (prior)','GP % change (pp)']
        hdr(ws, 5, heads); off = 1 if extra else 0; c0 = 2+off
        r0 = 6; r1 = r0+len(keys)-1; tr = r1+1
        def S(col, r, yr=None, ytd=False):
            e = ''
            if yr: e += f',{CA("J")},{yr}'
            if ytd: e += f',{CA("K")},"<="&Monthly_MIS!$B$6'
            return f'SUMIFS({CA(col)},{CA(key_col)},$A{r},{CA("C")},"Actual"{e})'
        for i, k in enumerate(keys):
            r = r0+i; put(ws, r, 1, k, font=fbb)
            if extra: put(ws, r, 2, extra[1][k])
            R, C_, G_, GM, SH, YC, YP, YY, GC, GP_, GD = [L(c0+j) for j in range(11)]
            put(ws, r, c0, '='+S('N', r), INR); put(ws, r, c0+1, '='+S('O', r), INR); put(ws, r, c0+2, f'={R}{r}-{C_}{r}', INR)
            put(ws, r, c0+3, f'=IF({R}{r}=0,"n/a",{G_}{r}/{R}{r})', PCT); put(ws, r, c0+4, f'={R}{r}/{R}${tr}', PCT)
            put(ws, r, c0+5, '='+S('N', r, 'Monthly_MIS!$B$5', True), INR); put(ws, r, c0+6, '='+S('N', r, '(Monthly_MIS!$B$5-1)', True), INR)
            put(ws, r, c0+7, f'=IF({YP}{r}=0,"n/a",{YC}{r}/{YP}{r}-1)', PCTS)
            gpc = f'({S("N", r, "Monthly_MIS!$B$5", True)}-{S("O", r, "Monthly_MIS!$B$5", True)})'; gpp = f'({S("N", r, "(Monthly_MIS!$B$5-1)", True)}-{S("O", r, "(Monthly_MIS!$B$5-1)", True)})'
            put(ws, r, c0+8, f'=IF({YC}{r}=0,"n/a",{gpc}/{YC}{r})', PCT); put(ws, r, c0+9, f'=IF({YP}{r}=0,"n/a",{gpp}/{YP}{r})', PCT)
            put(ws, r, c0+10, f'=IFERROR(({GC}{r}-{GP_}{r})*100,"n/a")', PP)
        put(ws, tr, 1, 'Total', font=fbb, fill=tfill)
        if extra: put(ws, tr, 2, None, fill=tfill)
        for j in (0,1,2,5,6): put(ws, tr, c0+j, f'=SUM({L(c0+j)}{r0}:{L(c0+j)}{r1})', INR, font=fbb, fill=tfill)
        put(ws, tr, c0+3, f'={L(c0+2)}{tr}/{L(c0)}{tr}', PCT, font=fbb, fill=tfill); put(ws, tr, c0+4, f'=SUM({L(c0+4)}{r0}:{L(c0+4)}{r1})', PCT, font=fbb, fill=tfill)
        put(ws, tr, c0+7, f'={L(c0+5)}{tr}/{L(c0+6)}{tr}-1', PCTS, font=fbb, fill=tfill)
        put(ws, tr, c0+8, f'=Monthly_MIS!H{PS+1}', PCT, font=fbb, fill=tfill); put(ws, tr, c0+9, f'=Monthly_MIS!H{PS+2}', PCT, font=fbb, fill=tfill)
        put(ws, tr, c0+10, f'=({L(c0+8)}{tr}-{L(c0+9)}{tr})*100', PP, font=fbb, fill=tfill)
        widths(ws, {'A':22, **{L(c):16 for c in range(2, 15)}}); ws.freeze_panes = 'B6'
        # monthly revenue matrix
        mr = tr+3; ws.cell(mr-1, 1, f'Monthly net revenue by {keylabel.lower()} (INR, actuals)').font = fbb
        hdr(ws, mr, [keylabel]+months+['Total'])
        for i, k in enumerate(keys):
            r = mr+1+i; put(ws, r, 1, k, font=fbb)
            for j, mo in enumerate(months):
                put(ws, r, 2+j, f'=SUMIFS({CA("N")},{CA(key_col)},$A{r},{CA("J")},VALUE(LEFT({L(2+j)}${mr},4)),{CA("K")},VALUE(RIGHT({L(2+j)}${mr},2)),{CA("C")},"Actual")', INR)
            put(ws, r, 2+len(months), f'=SUM(B{r}:{L(1+len(months))}{r})', INR, font=fbb)
        rt = mr+1+len(keys); put(ws, rt, 1, 'Total', font=fbb, fill=tfill)
        for j in range(len(months)+1): put(ws, rt, 2+j, f'=SUM({L(2+j)}{mr+1}:{L(2+j)}{rt-1})', INR, font=fbb, fill=tfill)
        for j in range(len(months)+1): ws.column_dimensions[L(2+j)].width = max(ws.column_dimensions[L(2+j)].width or 0, 14)
        # charts
        c1 = BarChart(); c1.type = 'col'; c1.title = f'Net Revenue & Gross Profit by {keylabel} (INR)'; c1.height = 8.5; c1.width = 17
        c1.add_data(Reference(ws, min_col=c0, min_row=5, max_row=r1), titles_from_data=True); c1.add_data(Reference(ws, min_col=c0+2, min_row=5, max_row=r1), titles_from_data=True)
        c1.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r1)); c1.series[0].graphicalProperties.solidFill = NAVY; c1.series[1].graphicalProperties.solidFill = TEAL
        c1.y_axis.number_format = '#,##0'; c1.x_axis.delete = False; c1.y_axis.delete = False; c1.legend.position = 'b'; c1.y_axis.majorGridlines = None
        c2 = BarChart(); c2.type = 'col'; c2.title = f'Gross Profit % by {keylabel}'; c2.height = 8.5; c2.width = 17
        c2.add_data(Reference(ws, min_col=c0+3, min_row=5, max_row=r1), titles_from_data=True); c2.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r1))
        c2.series[0].graphicalProperties.solidFill = ORANGE; c2.y_axis.number_format = '0%'; c2.y_axis.scaling.min = 0.25; c2.y_axis.scaling.max = 0.45
        c2.x_axis.delete = False; c2.y_axis.delete = False; c2.legend = None; c2.y_axis.majorGridlines = None
        ws.add_chart(c1, f'A{rt+3}'); ws.add_chart(c2, f'{L(7)}{rt+3}')
        return dict(r0=r0, r1=r1, tr=tr, c0=c0)

    regions = sorted(x['Customer_Master'].Region.dropna().unique())
    prods = list(x['Product_Master'].Product_Name); pcat = dict(zip(x['Product_Master'].Product_Name, x['Product_Master'].Category))
    RG = dim_sheet('Region_Performance', 'Region-wise Performance - Actuals', 'F', regions, 'Region')
    PR = dim_sheet('Product_Performance', 'Product-wise Performance - Actuals', 'I', prods, 'Product', extra=('Category', pcat))
    pw0 = W['Product_Performance']; put(pw0, 5, 14, 'Regions selling (actual)', font=fh, fill=hfill, al=Alignment(horizontal='center', vertical='center', wrap_text=True))
    for r_ in range(PR['r0'], PR['r1']+1):
        put(pw0, r_, 14, f'=SUMPRODUCT(--(COUNTIFS({CA("I")},$A{r_},{CA("F")},Region_Performance!$A${RG["r0"]}:$A${RG["r1"]},{CA("C")},"Actual")>0))', '0')
    pw0.column_dimensions['N'].width = 14
    # category roll-up on product sheet (placed to the right of charts area is awkward; place below charts)
    pw = W['Product_Performance']; cats = sorted(set(pcat.values())); cr0 = PR['tr']+3+len(prods)+3+22
    pw.cell(cr0-1, 1, 'Category roll-up (actuals)').font = fbb
    hdr(pw, cr0, ['Category','','Net Revenue (INR)','COGS (INR)','Gross Profit (INR)','GP %','Revenue share'])
    for i, cname in enumerate(cats):
        r = cr0+1+i; put(pw, r, 1, cname, font=fbb); put(pw, r, 2, None)
        put(pw, r, 3, f'=SUMIFS({CA("N")},{CA("H")},$A{r},{CA("C")},"Actual")', INR); put(pw, r, 4, f'=SUMIFS({CA("O")},{CA("H")},$A{r},{CA("C")},"Actual")', INR)
        put(pw, r, 5, f'=C{r}-D{r}', INR); put(pw, r, 6, f'=IF(C{r}=0,"n/a",E{r}/C{r})', PCT); put(pw, r, 7, f'=C{r}/C${cr0+1+len(cats)}', PCT)
    r = cr0+1+len(cats); put(pw, r, 1, 'Total', font=fbb, fill=tfill); put(pw, r, 2, None, fill=tfill)
    for c in (3,4,5): put(pw, r, c, f'=SUM({L(c)}{cr0+1}:{L(c)}{r-1})', INR, font=fbb, fill=tfill)
    put(pw, r, 6, f'=E{r}/C{r}', PCT, font=fbb, fill=tfill); put(pw, r, 7, f'=SUM(G{cr0+1}:G{r-1})', PCT, font=fbb, fill=tfill)
    CATROW = r


    # ---- Budget integrity + SYNTHETIC v2 (separate version; original Sales_Budget untouched)
    BVr0 = 13; extra_err = ''; BVC = None; BVA = None
    mrng = lambda c: f'Monthly_MIS!${c}${m0}:${c}${m1}'
    if bd:
        nB = raw_dim['Sales_Budget']; nBd = nB-1
        BVr1 = BVr0+nBd-1
        BVC = lambda c: f"Budget_v2_Synthetic!${c}${BVr0}:${c}${BVr1}"
        red = Font(name=F, size=12, bold=True, color='C00000')
        if bd['v2']:
            bw = W['Budget_v2_Synthetic']; bw.sheet_view.showGridLines = False; bw.sheet_properties.tabColor = 'C00000'
            bw['A1'] = 'SYNTHETIC CALIBRATED BUDGET - VERSION v2 - NOT AN APPROVED BUDGET'; bw['A1'].font = red
            bw['A2'] = 'Separate version for illustration only. The original Sales_Budget sheet is carried unchanged and is never overwritten. Requires finance approval before any use in reporting.'; bw['A2'].font = fs
            for rr_, lab, fm, fmt_, note in (
                (4, 'Calibration year (input)', bd['calib_year'], '0', 'Latest fully-actual year before the current year; avoids calibrating to the period being measured.'),
                (5, 'Actual net revenue in calibration year', f'=SUMIFS({CA("N")},{CA("J")},B4,{CA("C")},"Actual")', INR, 'All actual rows (Calc_Actuals).'),
                (6, 'Original budget revenue, calibration year, in-scope cells', f'=SUMIFS($D${BVr0}:$D${BVr1},$G${BVr0}:$G${BVr1},B4,$H${BVr0}:$H${BVr1},1)', INR, 'In-scope = Region x Category cell that has at least one actual sale in the source.'),
                (7, 'Calibration factor k = actual / original budget', '=IF(B6=0,0,B5/B6)', '0.0000', 'One uniform factor applied to every cell and month: preserves the original budget shape (regional/category pattern, year uplift, 38% margin).'),
                (8, 'Original budget GP % (unchanged in v2)', f'=SUM($F${BVr0}:$F${BVr1})/SUM($D${BVr0}:$D${BVr1})', PCT, 'Margin assumption is NOT calibrated to actuals.')):
                put(bw, rr_, 1, lab, font=fbb, fill=gfill); put(bw, rr_, 2, fm, fmt_, font=(fin if rr_ == 4 else fb)); bw.cell(rr_, 3, note).font = fs
            bw.merge_cells('A9:O9'); bw['A9'] = ('ASSUMPTIONS: (1) Business scale of the budget is unknown; v2 assumes the budget should be the same order of magnitude as FY' + str(bd['calib_year']) + ' actuals. '
                '(2) Budget cells with no actual sales anywhere in the source are retained but flagged out-of-scope and excluded from variance. (3) FY' + str(bd['calib_year']) + ' variance is ~0 by construction (circular); only periods after it are out-of-sample. '
                '(4) A single k does not fix the non-constant ratio between cells, so cell-level variances remain indicative only.'); bw['A9'].alignment = wrap; bw['A9'].font = fb; bw.row_dimensions[9].height = 62
            hdr(bw, 12, ['Month','Region','Category','Orig Revenue (INR)','Orig COGS (INR)','Orig GP (INR)','Year','In scope (1/0)','Calendar status','v2 Revenue (INR)','v2 COGS (INR)','v2 GP (INR)','Actual Revenue (INR)','Actual GP (INR)','Month no.','Key'])
            for i in range(nBd):
                rr = 2+i; r = BVr0+i
                fm = {1:f'=Sales_Budget!A{rr}', 2:f'=Sales_Budget!B{rr}', 3:f'=Sales_Budget!C{rr}', 4:f'=Sales_Budget!E{rr}', 5:f'=Sales_Budget!F{rr}', 6:f'=Sales_Budget!G{rr}',
                      7:f'=VALUE(LEFT(A{r},4))', 8:f'=IF(COUNTIFS({CA("F")},B{r},{CA("H")},C{r},{CA("C")},"Actual")>0,1,0)', 9:f'=IFERROR(INDEX({CAL("E")},MATCH(A{r},{CAL("A")},0)),"UNMATCHED")',
                      10:f'=D{r}*$B$7', 11:f'=E{r}*$B$7', 12:f'=F{r}*$B$7',
                      13:f'=IF(I{r}="Actual",SUMIFS({CA("N")},{CA("J")},G{r},{CA("K")},O{r},{CA("F")},B{r},{CA("H")},C{r},{CA("C")},"Actual"),0)',
                      14:f'=IF(I{r}="Actual",SUMIFS({CA("P")},{CA("J")},G{r},{CA("K")},O{r},{CA("F")},B{r},{CA("H")},C{r},{CA("C")},"Actual"),0)',
                      15:f'=VALUE(RIGHT(A{r},2))', 16:f'=A{r}&"|"&B{r}&"|"&C{r}'}
                for c, v in fm.items():
                    cell = bw.cell(r, c, v); cell.font = fb
                    if c in (4,5,6,10,11,12,13,14): cell.number_format = INR
            widths(bw, {'A':44,'B':18,'C':16,'D':18,'E':16,'F':16,'G':8,'H':10,'I':12,'J':16,'K':16,'L':16,'M':16,'N':16,'O':9,'P':26}); bw.freeze_panes = f'D{BVr0}'
            extra_err = f"+SUMPRODUCT(--ISERROR(Budget_v2_Synthetic!$A${BVr0}:$P${BVr1}))"
            # ---------------- Budget_vs_Actual_v2
            ba = W['Budget_vs_Actual_v2']; ba.sheet_view.showGridLines = False; ba.sheet_properties.tabColor = 'C00000'
            ba['A1'] = 'ILLUSTRATIVE ONLY - Variance vs SYNTHETIC calibrated budget v2 (unapproved)'; ba['A1'].font = red
            ba['A2'] = 'In-scope Region x Category cells only. The original Sales_Budget is not used here. Actuals are the same actual-only figures as Monthly_MIS. See Budget_Integrity for why.'; ba['A2'].font = fs
            hdr(ba, 4, ['Month','v2 Budget Revenue','Actual Revenue','Variance (INR)','Variance %','v2 Budget GP','Actual GP','GP Variance (INR)','Basis'])
            b0 = 5
            for i, mo in enumerate(months):
                r = b0+i; yk = f'{BVC("G")},VALUE(LEFT($A{r},4)),{BVC("O")},VALUE(RIGHT($A{r},2)),{BVC("H")},1'
                put(ba, r, 1, mo); put(ba, r, 2, f'=SUMIFS({BVC("J")},{yk})', INR); put(ba, r, 3, f'=SUMIFS({BVC("M")},{yk})', INR)
                put(ba, r, 4, f'=C{r}-B{r}', '#,##0;-#,##0;0'); put(ba, r, 5, f'=IF(B{r}=0,"n/a",D{r}/B{r})', PCTS)
                put(ba, r, 6, f'=SUMIFS({BVC("L")},{yk})', INR); put(ba, r, 7, f'=SUMIFS({BVC("N")},{yk})', INR); put(ba, r, 8, f'=G{r}-F{r}', '#,##0;-#,##0;0')
                put(ba, r, 9, f'=IF(VALUE(LEFT(A{r},4))=Budget_v2_Synthetic!$B$4,"Calibration year (circular)","Out-of-sample")')
            b1 = b0+len(months)-1; bt = b1+1
            put(ba, bt, 1, 'Total', font=fbb, fill=tfill)
            for c in (2,3,6,7): put(ba, bt, c, f'=SUM({L(c)}{b0}:{L(c)}{b1})', INR, font=fbb, fill=tfill)
            put(ba, bt, 4, f'=C{bt}-B{bt}', '#,##0;-#,##0;0', font=fbb, fill=tfill); put(ba, bt, 5, f'=IF(B{bt}=0,"n/a",D{bt}/B{bt})', PCTS, font=fbb, fill=tfill)
            put(ba, bt, 8, f'=G{bt}-F{bt}', '#,##0;-#,##0;0', font=fbb, fill=tfill); put(ba, bt, 9, None, fill=tfill)
            pr0 = bt+3; ba.cell(pr0-1, 1, 'Period view').font = fbb
            hdr(ba, pr0, ['Period','v2 Budget Revenue','Actual Revenue','Variance (INR)','Variance %','Reading'])
            put(ba, pr0+1, 1, '="FY"&Budget_v2_Synthetic!B4&" (calibration year)"', font=fbb)
            put(ba, pr0+1, 2, f'=SUMIFS({BVC("J")},{BVC("G")},Budget_v2_Synthetic!$B$4,{BVC("H")},1)', INR); put(ba, pr0+1, 3, f'=SUMIFS({BVC("M")},{BVC("G")},Budget_v2_Synthetic!$B$4,{BVC("H")},1)', INR)
            put(ba, pr0+1, 6, 'Variance ~0 by construction - not informative', font=fs)
            put(ba, pr0+2, 1, '="YTD "&Monthly_MIS!B5&" (out-of-sample)"', font=fbb)
            put(ba, pr0+2, 2, f'=SUMIFS({BVC("J")},{BVC("G")},Monthly_MIS!$B$5,{BVC("O")},"<="&Monthly_MIS!$B$6,{BVC("I")},"Actual",{BVC("H")},1)', INR)
            put(ba, pr0+2, 3, f'=SUMIFS({BVC("M")},{BVC("G")},Monthly_MIS!$B$5,{BVC("O")},"<="&Monthly_MIS!$B$6,{BVC("H")},1)', INR)
            put(ba, pr0+2, 6, 'Only meaningful comparison; still synthetic and unapproved', font=fs)
            for r in (pr0+1, pr0+2):
                put(ba, r, 4, f'=C{r}-B{r}', '#,##0;-#,##0;0'); put(ba, r, 5, f'=IF(B{r}=0,"n/a",D{r}/B{r})', PCTS)
            YTDB = pr0+2
            rc0 = pr0+6; ba.cell(rc0-1, 1, '="YTD "&Monthly_MIS!B5&" by Region x Category"').font = fbb
            hdr(ba, rc0, ['Region','Category','In scope?','YTD v2 Budget Revenue','YTD Actual Revenue','Variance (INR)','Variance %','Note'])
            for i, (rg, ct) in enumerate(bd['cells']):
                r = rc0+1+i; put(ba, r, 1, rg); put(ba, r, 2, ct)
                put(ba, r, 3, f'=IF(COUNTIFS({CA("F")},A{r},{CA("H")},B{r},{CA("C")},"Actual")>0,"Yes","No")', al=Alignment(horizontal='center'))
                cr_ = f'{BVC("B")},$A{r},{BVC("C")},$B{r},{BVC("G")},Monthly_MIS!$B$5,{BVC("O")},"<="&Monthly_MIS!$B$6,{BVC("I")},"Actual"'
                put(ba, r, 4, f'=SUMIFS({BVC("J")},{cr_})', INR); put(ba, r, 5, f'=SUMIFS({BVC("M")},{cr_})', INR)
                put(ba, r, 6, f'=IF(C{r}="Yes",E{r}-D{r},"excluded")', '#,##0;-#,##0;0', al=Alignment(horizontal='right')); put(ba, r, 7, f'=IF(C{r}="Yes",IF(D{r}=0,"n/a",F{r}/D{r}),"excluded")', PCTS, al=Alignment(horizontal='right'))
                put(ba, r, 8, f'=IF(C{r}="Yes","","No actual sales for this cell in the source - budget excluded from variance")', font=fs)
            re_ = rc0+len(bd['cells']); rt = re_+1
            put(ba, rt, 1, 'In-scope total', font=fbb, fill=tfill); put(ba, rt, 2, None, fill=tfill); put(ba, rt, 3, None, fill=tfill)
            put(ba, rt, 4, f'=SUMIFS(D{rc0+1}:D{re_},C{rc0+1}:C{re_},"Yes")', INR, font=fbb, fill=tfill); put(ba, rt, 5, f'=SUMIFS(E{rc0+1}:E{re_},C{rc0+1}:C{re_},"Yes")', INR, font=fbb, fill=tfill)
            put(ba, rt, 6, f'=E{rt}-D{rt}', '#,##0;-#,##0;0', font=fbb, fill=tfill); put(ba, rt, 7, f'=IF(D{rt}=0,"n/a",F{rt}/D{rt})', PCTS, font=fbb, fill=tfill); put(ba, rt, 8, None, fill=tfill)
            put(ba, rt+1, 1, 'Excluded (no-sales cells) v2 budget', font=fbb); put(ba, rt+1, 4, f'=SUMIFS(D{rc0+1}:D{re_},C{rc0+1}:C{re_},"No")', INR)
            widths(ba, {'A':30,'B':20,'C':18,'D':22,'E':20,'F':18,'G':14,'H':54,'I':26}); ba.freeze_panes = 'A5'
            c3 = BarChart(); c3.type = 'col'; c3.title = 'Actual vs SYNTHETIC v2 Budget - Revenue (illustrative)'; c3.height = 8.5; c3.width = 24
            c3.add_data(Reference(ba, min_col=2, min_row=4, max_row=b1), titles_from_data=True); c3.add_data(Reference(ba, min_col=3, min_row=4, max_row=b1), titles_from_data=True)
            c3.set_categories(Reference(ba, min_col=1, min_row=b0, max_row=b1)); c3.series[0].graphicalProperties.solidFill = 'A6A6A6'; c3.series[1].graphicalProperties.solidFill = NAVY
            c3.x_axis.delete = False; c3.y_axis.delete = False; c3.y_axis.number_format = '#,##0'; c3.legend.position = 'b'; c3.y_axis.majorGridlines = None
            ba.add_chart(c3, f'A{rt+4}')
            BVA = dict(bt=bt, YTDB=YTDB, rt=rt)
        # ---------------- Budget_Integrity (always built when a budget exists)
        bi = W['Budget_Integrity']; bi.sheet_view.showGridLines = False; bi.sheet_properties.tabColor = 'C00000'
        bi['A1'] = 'Budget Integrity Review - Sales_Budget vs actuals'; bi['A1'].font = ft
        bi['A2'] = 'Original Sales_Budget is untouched. Findings are computed from the source at build time (static text); live cross-check cells below recompute the headline ratios from the workbook.'; bi['A2'].font = fs
        put(bi, 4, 1, 'Live cross-check', font=fh, fill=hfill); put(bi, 4, 2, 'Value', font=fh, fill=hfill)
        put(bi, 5, 1, 'Budget / actual revenue - all 12 Region x Category cells', font=fbb, fill=gfill)
        put(bi, 5, 2, f'=SUMPRODUCT(ISNUMBER(MATCH(Sales_Budget!$A$2:$A${nB},{mrng("A")},0))*Sales_Budget!$E$2:$E${nB})/Monthly_MIS!E{mt}', '0.00"x"')
        if bd['v2']:
            put(bi, 6, 1, 'Budget / actual revenue - only cells that have sales', font=fbb, fill=gfill)
            put(bi, 6, 2, f'=SUMPRODUCT(ISNUMBER(MATCH({BVC("A")},{mrng("A")},0))*{BVC("H")}*{BVC("D")})/Monthly_MIS!E{mt}', '0.00"x"')
            put(bi, 7, 1, 'Share of budget in cells with no sales', font=fbb, fill=gfill); put(bi, 7, 2, '=1-B6/B5', PCT)
        else:
            put(bi, 6, 1, 'v2 not built', font=fbb, fill=gfill); put(bi, 6, 2, 'no calibration year or budget not flagged', font=fs)
        hdr(bi, 9, ['Test','Method','Finding','Verdict'])
        for i, (t1, t2, t3, t4) in enumerate(bd['tests']):
            r = 10+i; put(bi, r, 1, t1, font=fbb, al=wrap); put(bi, r, 2, t2, al=wrap); put(bi, r, 3, t3, al=wrap); put(bi, r, 4, t4, font=fbb, al=wrap); bi.row_dimensions[r].height = 78 if len(t3) > 190 else 52
        r = 10+len(bd['tests'])+1
        bi.cell(r, 1, 'Root-cause conclusion').font = fbb; bi.merge_cells(start_row=r+1, start_column=1, end_row=r+1, end_column=4); c = bi.cell(r+1, 1, bd['root']); c.alignment = wrap; c.font = fb; bi.row_dimensions[r+1].height = 92
        bi.cell(r+3, 1, 'Controlled correction proposed').font = fbb
        txt = (f"Budget_v2_Synthetic applies one uniform factor k = {bd['k']:.4f} to the original budget (k = FY{bd['calib_year']} actual revenue / original FY{bd['calib_year']} budget in cells that have sales). Shape, 38% margin and year uplift are unchanged; no-sales cells are excluded from variance. "
               "It is SYNTHETIC and UNAPPROVED. It does not resolve the cause; it only makes a like-for-like illustration possible. Human decision needed: confirm the intended budget scale, or supply a corrected budget." if bd['v2'] else 'No synthetic version built (budget not flagged, or no full prior actual year).')
        bi.merge_cells(start_row=r+4, start_column=1, end_row=r+4, end_column=4); c = bi.cell(r+4, 1, txt); c.alignment = wrap; c.font = fb; bi.row_dimensions[r+4].height = 80
        widths(bi, {'A':54,'B':44,'C':92,'D':30})

    # ---- Control_Checks
    cw = W['Control_Checks']; title(cw, 'Control Checks - Reconciliations & Data Integrity', 'Live formulas. Blue = hard-coded expected values taken from the source file\'s own Control_Checks sheet. Tolerance for INR sums = 0.01.')
    put(cw, 3, 1, 'Overall status', font=fbb, fill=gfill); cw.merge_cells('A3:B3')
    hdr(cw, 5, ['#','Group','Control','Method / calculation','Expected','Actual','Difference','Tolerance','Result','Exceptions'])
    mrng = lambda c: f'Monthly_MIS!${c}${m0}:${c}${m1}'
    b = lambda k: base.get(k)
    rawsum = lambda col, st: f'SUMIFS({ST(col)},{ST("C")},"{st}")'
    chk = []
    def add(group, name, method, exp, act, tol=0, blue=False): chk.append((group, name, method, exp, act, tol, blue))
    n_rows = f'ROWS({ST("A")})'
    add('Record counts','Sales_Transactions rows','Expected: source Control_Checks baseline. Actual: COUNTA of Invoice_ID.', b('Sales transactions'), f'=COUNTA({ST("A")})', blue=True)
    add('Record counts','Actual-status rows','Expected: source baseline. Actual: COUNTIF Status="Actual".', b('Actual sales rows'), f'=COUNTIF({ST("C")},"Actual")', blue=True)
    add('Record counts','Forecast-status rows','Expected: source baseline. Actual: COUNTIF Status="Forecast".', b('Forecast sales rows'), f'=COUNTIF({ST("C")},"Forecast")', blue=True)
    add('Record counts','Rows with Status other than Actual/Forecast','Expected 0. Actual: rows - Actual - Forecast.', 0, f'={n_rows}-COUNTIF({ST("C")},"Actual")-COUNTIF({ST("C")},"Forecast")')
    add('Record counts','Blank cells in Sales_Transactions','Expected 0. COUNTBLANK over A:K data range.', 0, f'=COUNTBLANK(Sales_Transactions!$A$2:$K${nS})')
    add('Keys','Unique Invoice_IDs','Expected: source baseline. Actual: SUMPRODUCT(1/COUNTIF) distinct count.', b('Unique invoice IDs'), f'=ROUND(SUMPRODUCT(1/COUNTIF({ST("A")},{ST("A")})),0)', blue=True)
    add('Keys','Duplicate Invoice_ID rows','Expected 0. Rows - distinct Invoice_IDs.', 0, f'={n_rows}-ROUND(SUMPRODUCT(1/COUNTIF({ST("A")},{ST("A")})),0)')
    add('Keys','Unmatched Product_ID (Sales -> Product_Master)','Expected: source baseline. Rows whose Product_ID is absent from Product_Master.', b('Unmatched products'), f'=SUMPRODUCT(--ISNA(MATCH({ST("E")},{PM("A")},0)))', blue=True)
    add('Keys','Unmatched Customer_ID (Sales -> Customer_Master)','Expected: source baseline. Rows whose Customer_ID is absent from Customer_Master.', b('Unmatched customers'), f'=SUMPRODUCT(--ISNA(MATCH({ST("D")},{CM("A")},0)))', blue=True)
    add('Keys','Unmatched Month (Sales -> Calendar)','Expected 0. Rows whose Month is absent from Calendar.', 0, f'=SUMPRODUCT(--ISNA(MATCH({ST("B")},{CAL("A")},0)))')
    add('Keys','Duplicate Product_ID in Product_Master','Expected 0. Rows - distinct IDs.', 0, f'=ROWS({PM("A")})-ROUND(SUMPRODUCT(1/COUNTIF({PM("A")},{PM("A")})),0)')
    add('Keys','Duplicate Customer_ID in Customer_Master','Expected 0. Rows - distinct IDs.', 0, f'=ROWS({CM("A")})-ROUND(SUMPRODUCT(1/COUNTIF({CM("A")},{CM("A")})),0)')
    add('Status','Transaction Status vs Calendar Reporting_Status mismatches','Expected 0. Count of rows where Status <> Calendar.Reporting_Status (Calc_Actuals col V).', 0, f'=SUM({CA("V")})')
    add('Row arithmetic','Net_Revenue <> Units x Net_Price (>0.01)','Expected 0 rows. SUMPRODUCT of ABS(F*G-I)>0.01.', 0, f'=SUMPRODUCT(--(ABS({ST("F")}*{ST("G")}-{ST("I")})>0.01))')
    add('Row arithmetic','COGS <> Units x Unit_Cost (>0.01)','Expected 0 rows. SUMPRODUCT of ABS(F*H-J)>0.01.', 0, f'=SUMPRODUCT(--(ABS({ST("F")}*{ST("H")}-{ST("J")})>0.01))')
    add('Row arithmetic','Gross_Profit <> Net_Revenue - COGS (>0.01)','Expected 0 rows. SUMPRODUCT of ABS(I-J-K)>0.01.', 0, f'=SUMPRODUCT(--(ABS({ST("I")}-{ST("J")}-{ST("K")})>0.01))')
    add('Row arithmetic','Non-positive Units / Net_Price / Unit_Cost','Expected 0. COUNTIF <=0 on the three columns.', 0, f'=COUNTIF({ST("F")},"<=0")+COUNTIF({ST("G")},"<=0")+COUNTIF({ST("H")},"<=0")')
    for lab, mcol, rcol, bk in (('Net revenue','E','I','Actual net revenue INR'), ('COGS','F','J','Actual COGS INR'), ('Gross profit','G','K','Actual gross profit INR')):
        add('Reconciliation', f'{lab}: Monthly_MIS total vs raw source (Status=Actual)', f'Expected: SUMIFS on raw Sales_Transactions col {rcol}. Actual: Monthly_MIS total row.', '='+rawsum(rcol,'Actual'), f'=Monthly_MIS!{mcol}{mt}', 0.01)
        add('Reconciliation', f'{lab}: Monthly_MIS total vs source-file baseline', 'Expected: baseline in source Control_Checks sheet (hard-coded). Actual: Monthly_MIS total row.', b(bk), f'=Monthly_MIS!{mcol}{mt}', 0.01, True)
    for lab, mcol, rcol in (('Net revenue','E','I'), ('COGS','F','J')):
        add('Reconciliation', f'{lab}: months where Monthly_MIS differs from raw SUMIFS by >0.01', 'Expected 0 months. Month-by-month comparison to Monthly_MIS cols L/M (SUMPRODUCT directly on raw Sales_Transactions).', 0,
            f'=SUMPRODUCT(--(ABS({mrng(mcol)}-{mrng("L" if mcol=="E" else "M")})>0.01))')
    rc = RG['c0']; pc = PR['c0']
    for lab, off, mcol in (('Net revenue',0,'E'), ('COGS',1,'F'), ('Gross profit',2,'G')):
        add('Reconciliation', f'{lab}: Region_Performance total vs Monthly_MIS', 'Expected: Monthly_MIS total. Actual: sum of regions.', f'=Monthly_MIS!{mcol}{mt}', f"=Region_Performance!{L(rc+off)}{RG['tr']}", 0.01)
        add('Reconciliation', f'{lab}: Product_Performance total vs Monthly_MIS', 'Expected: Monthly_MIS total. Actual: sum of products.', f'=Monthly_MIS!{mcol}{mt}', f"=Product_Performance!{L(pc+off)}{PR['tr']}", 0.01)
    for lab, off, mrow in (('YTD current revenue',5,PS+1), ('YTD prior-year revenue',6,PS+2)):
        add('Reconciliation', f'{lab}: Region_Performance total vs Monthly_MIS period summary', 'Expected: Monthly_MIS period summary (SUMIFS by Year/Month_No). Actual: sum of regions.', f'=Monthly_MIS!E{mrow}', f"=SUM(Region_Performance!{L(rc+off)}{RG['r0']}:{L(rc+off)}{RG['r1']})", 0.01)
        add('Reconciliation', f'{lab}: Product_Performance total vs Monthly_MIS period summary', 'Expected: Monthly_MIS period summary. Actual: sum of products.', f'=Monthly_MIS!E{mrow}', f"=SUM(Product_Performance!{L(pc+off)}{PR['r0']}:{L(pc+off)}{PR['r1']})", 0.01)
    add('Reconciliation','YTD gross profit (current): sum of region GP% x YTD revenue vs Monthly_MIS','Expected: Monthly_MIS YTD gross profit. Actual: SUMPRODUCT(region YTD GP% x YTD revenue).', f'=Monthly_MIS!G{PS+1}', f"=SUMPRODUCT(Region_Performance!{L(rc+8)}{RG['r0']}:{L(rc+8)}{RG['r1']},Region_Performance!{L(rc+5)}{RG['r0']}:{L(rc+5)}{RG['r1']})", 0.01)
    add('Reconciliation','YTD gross profit (current): sum of product GP% x YTD revenue vs Monthly_MIS','Expected: Monthly_MIS YTD gross profit. Actual: SUMPRODUCT(product YTD GP% x YTD revenue).', f'=Monthly_MIS!G{PS+1}', f"=SUMPRODUCT(Product_Performance!{L(pc+8)}{PR['r0']}:{L(pc+8)}{PR['r1']},Product_Performance!{L(pc+5)}{PR['r0']}:{L(pc+5)}{PR['r1']})", 0.01)
    add('Reconciliation','Net revenue: Category roll-up vs Monthly_MIS','Expected: Monthly_MIS total. Actual: sum of categories.', f'=Monthly_MIS!E{mt}', f'=Product_Performance!C{CATROW}', 0.01)
    add('Reconciliation','Gross profit identity on Monthly_MIS','Expected: sum of raw Gross_Profit_INR (Actual). Actual: Revenue - COGS from Monthly_MIS total.', '='+rawsum('K','Actual'), f'=Monthly_MIS!E{mt}-Monthly_MIS!F{mt}', 0.01)
    add('Scope','Forecast excluded: Monthly_MIS + Forecast revenue = all source revenue','Expected: SUM of all raw Net_Revenue_INR. Actual: Monthly total + Forecast-status revenue.', f'=SUM({ST("I")})', f'=Monthly_MIS!E{mt}+'+rawsum('I','Forecast'), 0.01)
    add('Scope','Forecast months appearing in Monthly_MIS','Expected 0. Calendar months flagged Forecast that are listed in Monthly_MIS.', 0, f'=SUMPRODUCT(({CAL("E")}="Forecast")*ISNUMBER(MATCH({CAL("A")},{mrng("A")},0)))')
    add('Scope','Calendar Actual months vs months listed in Monthly_MIS','Expected: count of Actual months in Calendar. Actual: months listed.', f'=COUNTIF({CAL("E")},"Actual")', f'=COUNTA({mrng("A")})')
    add('Scope','Listed months with zero invoices','Expected 0.', 0, f'=COUNTIF({mrng("C")},0)')
    if bd and bd['v2']:
        add('Budget v2','Original Sales_Budget carried unchanged into Budget_v2 (revenue)','Expected: SUM of Sales_Budget revenue. Actual: SUM of original-revenue column in Budget_v2_Synthetic.', f'=SUM(Sales_Budget!$E$2:$E${nB})', f'=SUM({BVC("D")})', 0.01)
        add('Budget v2','Calibration identity: in-scope v2 budget = actual in calibration year','Expected: actual revenue, calibration year. Actual: SUM of in-scope v2 budget, calibration year.', '=Budget_v2_Synthetic!B5', f'=SUMIFS({BVC("J")},{BVC("G")},Budget_v2_Synthetic!B4,{BVC("H")},1)', 0.01)
        add('Budget v2','v2 gross profit = v2 revenue - v2 COGS','Expected: SUM v2 revenue - SUM v2 COGS. Actual: SUM v2 GP.', f'=SUM({BVC("J")})-SUM({BVC("K")})', f'=SUM({BVC("L")})', 0.01)
        add('Budget v2','Variance sheet actuals tie to Monthly_MIS (all actual sales in scope)','Expected: Monthly_MIS total revenue. Actual: Budget_vs_Actual_v2 total actual.', f'=Monthly_MIS!E{mt}', f'=Budget_vs_Actual_v2!C{BVA["bt"]}', 0.01)
        add('Budget v2','Budget months not found in Calendar','Expected 0.', 0, f'=COUNTIF({BVC("I")},"UNMATCHED")')
        add('Budget v2','Duplicate Month|Region|Category keys in budget','Expected 0.', 0, f'=ROWS({BVC("P")})-ROUND(SUMPRODUCT(1/COUNTIF({BVC("P")},{BVC("P")})),0)')
    add('Integrity','Formula errors in Calc_Actuals / report sheets','Expected 0. ISERROR count across Calc_Actuals and report tables.', 0,
        f"=SUMPRODUCT(--ISERROR(Calc_Actuals!$A$2:$W${nS}))+SUMPRODUCT(--ISERROR(Monthly_MIS!$A${m0}:$M${mt}))+SUMPRODUCT(--ISERROR(Region_Performance!$A${RG['r0']}:$N${RG['tr']}))+SUMPRODUCT(--ISERROR(Product_Performance!$A${PR['r0']}:$N${PR['tr']})){extra_err}")
    add('Integrity','Rows flagged in Calc_Actuals (any exception)','Expected 0. Filter Calc_Actuals col W = 1 to list exception invoices.', 0, f'=SUM({CA("W")})')
    r = 6
    for i, (g, nme, meth, exp, act, tol, blue) in enumerate(chk, 1):
        put(cw, r, 1, i); put(cw, r, 2, g); put(cw, r, 3, nme, al=wrap); put(cw, r, 4, meth, al=wrap)
        if exp is None:
            put(cw, r, 5, 'n/a (no baseline)', font=fs); put(cw, r, 6, act, '#,##0.00'); put(cw, r, 7, None); put(cw, r, 8, tol); put(cw, r, 9, 'N/A')
        else:
            e = put(cw, r, 5, exp, '#,##0.00' if (isinstance(exp, str) or tol) else '#,##0', font=fin if blue else fb)
            put(cw, r, 6, act, '#,##0.00' if tol else '#,##0'); put(cw, r, 7, f'=F{r}-E{r}', '#,##0.00;-#,##0.00;0.00' if tol else '#,##0;-#,##0;0'); put(cw, r, 8, tol, '0.00')
            put(cw, r, 9, f'=IF(ISNUMBER(G{r}),IF(ABS(G{r})<=H{r},"PASS","FAIL"),"FAIL")', font=fbb, al=Alignment(horizontal='center'))
        put(cw, r, 10, f'=IF(I{r}="PASS","None",IF(I{r}="N/A","No baseline in source","EXCEPTION: difference "&TEXT(G{r},"#,##0.00")&" - investigate; see Calc_Actuals filter on col W"))', al=wrap)
        cw.row_dimensions[r].height = 30
        r += 1
    last = r-1
    put(cw, 3, 3, f'=IF(COUNTIF(I6:I{last},"FAIL")=0,"ALL "&COUNTIF(I6:I{last},"PASS")&" CHECKS PASS",COUNTIF(I6:I{last},"FAIL")&" CHECK(S) FAIL")', font=fbb, fill=gfill)
    cw.conditional_formatting.add(f'I6:I{last}', CellIsRule(operator='equal', formula=['"PASS"'], fill=PatternFill('solid', bgColor='C6EFCE'), font=Font(color='006100', bold=True)))
    cw.conditional_formatting.add(f'I6:I{last}', CellIsRule(operator='equal', formula=['"FAIL"'], fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006', bold=True)))
    cw.conditional_formatting.add('C3', FormulaRule(formula=['ISNUMBER(SEARCH("FAIL",C3))'], fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006', bold=True)))
    cw.conditional_formatting.add('C3', FormulaRule(formula=['ISNUMBER(SEARCH("PASS",C3))'], fill=PatternFill('solid', bgColor='C6EFCE'), font=Font(color='006100', bold=True)))
    # observations
    r += 1; cw.cell(r, 1, 'Observations (information only - not pass/fail)').font = fbb; r += 1
    hdr(cw, r, ['#','Area','Observation','Calculation','Value','','','','Status','Note']); r += 1
    obs_rows = []
    if 'Sales_Budget' in x:
        nB = raw_dim['Sales_Budget']
        put(cw, r, 1, 'O1'); put(cw, r, 2, 'Budget scale'); put(cw, r, 3, 'Sales_Budget revenue for the same actual months vs actual net revenue (ratio)', al=wrap)
        put(cw, r, 4, 'Sales_Budget revenue for listed actual months / Monthly_MIS total revenue (all 12 Region x Category cells)', al=wrap)
        put(cw, r, 5, f'=SUMPRODUCT(ISNUMBER(MATCH(Sales_Budget!$A$2:$A${nB},{mrng("A")},0))*Sales_Budget!$E$2:$E${nB})/Monthly_MIS!E{mt}', '0.0"x"')
        for c in (6,7,8): put(cw, r, c, None)
        put(cw, r, 9, f'=IF(OR(E{r}>1.5,E{r}<0.67),"FLAG","OK")', font=fbb, al=Alignment(horizontal='center'))
        put(cw, r, 10, 'Budget is a separate scenario and is NOT compared with actuals. Scale mismatch needs business confirmation before any variance reporting.', al=wrap); cw.row_dimensions[r].height = 42; r += 1
    put(cw, r, 1, 'O2'); put(cw, r, 2, 'Scope'); put(cw, r, 3, 'Operating_Expenses, Workforce, Inventory, Receivables sheets not used in Day-1 MIS', al=wrap); put(cw, r, 4, 'Outside Day-1 scope (sales & gross profit). Not copied; original file untouched.', al=wrap)
    for c in range(5, 9): put(cw, r, c, None)
    put(cw, r, 9, 'INFO', font=fbb, al=Alignment(horizontal='center')); put(cw, r, 10, 'In the source, Operating_Expenses.Actual_INR is blank on all 75 Forecast rows (expected for forecasts).', al=wrap); cw.row_dimensions[r].height = 30; r += 1
    put(cw, r, 1, 'O3'); put(cw, r, 2, 'Region vs product'); put(cw, r, 3, 'Maximum number of regions in which any single product sold (actual)', al=wrap); put(cw, r, 4, 'MAX of Product_Performance col N. If 1, every product sells in one region only, so region and product effects cannot be separated.', al=wrap)
    put(cw, r, 5, f"=MAX(Product_Performance!N{PR['r0']}:N{PR['r1']})", '0')
    for c in (6,7,8): put(cw, r, c, None)
    put(cw, r, 9, f'=IF(E{r}<=1,"FLAG","OK")', font=fbb, al=Alignment(horizontal='center')); put(cw, r, 10, 'Region differences in revenue growth and margin are product-mix effects in this dataset; do not read them as regional performance without more data.', al=wrap); cw.row_dimensions[r].height = 48; r += 2
    # static exception snapshot via pandas (independent engine)
    s = x['Sales_Transactions']
    ex = []
    d = s[s.Invoice_ID.duplicated(keep=False)]; ex += [f'Duplicate Invoice_ID: {i}' for i in d.Invoice_ID.unique()]
    ex += [f'Unmatched Product_ID: {i}' for i in s[~s.Product_ID.isin(x['Product_Master'].Product_ID)].Invoice_ID]
    ex += [f'Unmatched Customer_ID: {i}' for i in s[~s.Customer_ID.isin(x['Customer_Master'].Customer_ID)].Invoice_ID]
    ex += [f'Arithmetic mismatch: {i}' for i in s[((s.Units*s.Net_Price_INR-s.Net_Revenue_INR).abs()>0.01)|((s.Units*s.Unit_Cost_INR-s.COGS_INR).abs()>0.01)|((s.Net_Revenue_INR-s.COGS_INR-s.Gross_Profit_INR).abs()>0.01)].Invoice_ID]
    mg = s.merge(x['Calendar'][['Month','Reporting_Status']], on='Month', how='left'); ex += [f'Status vs Calendar mismatch: {i}' for i in mg[mg.Status!=mg.Reporting_Status].Invoice_ID]
    cw.cell(r, 1, 'Exception detail - build-time snapshot from an independent pandas scan (static; the live counts above are authoritative)').font = fbb; r += 1
    if not ex: put(cw, r, 1, '-', border=False); cw.cell(r, 3, 'No exceptions found at build time.').font = fb
    for e in ex[:200]: cw.cell(r, 3, e).font = fb; r += 1
    widths(cw, {'A':5,'B':15,'C':44,'D':52,'E':18,'F':18,'G':14,'H':10,'I':10,'J':46}); cw.freeze_panes = 'A6'

    # ---- Executive_Summary
    es = W['Executive_Summary']; title(es, 'Nova Consumer Products - CFO Monthly Sales MIS', None)
    es['A2'] = '="Actuals "&Monthly_MIS!A9&" to "&Monthly_MIS!B4&"  |  INR  |  Synthetic data  |  Controls: "&Control_Checks!C3'; es['A2'].font = fs
    hdr(es, 4, ['Headline KPI','Value','Comment'])
    k = [('Net revenue (all actual months)', f'=Monthly_MIS!E{mt}', INR, 'Actual-status invoices only'), ('COGS', f'=Monthly_MIS!F{mt}', INR, ''), ('Gross profit', f'=Monthly_MIS!G{mt}', INR, ''),
         ('Gross profit %', f'=Monthly_MIS!H{mt}', PCT, ''), ('="YTD revenue "&Monthly_MIS!B5', f'=Monthly_MIS!E{PS+1}', INR, 'Jan to latest actual month'),
         ('YTD revenue change vs prior-year YTD', f'=Monthly_MIS!E{PS+4}', PCTS, 'Like-for-like months'), ('YTD GP % change vs prior-year YTD', f'=Monthly_MIS!H{PS+4}', PP, 'Percentage points'),
         ('Control checks', '=Control_Checks!C3', None, 'See Control_Checks')]
    for i, (a, v, fmt, cm) in enumerate(k):
        r = 5+i; put(es, r, 1, a, font=fbb); put(es, r, 2, v, fmt, al=Alignment(horizontal='right')); put(es, r, 3, cm, font=fs)
    pack = fact_pack(x, fx, bd, ef)
    obs, unv = commentary(x, fx, bd, ef)
    if commentary_override:
        ov = commentary_override.strip(); probs = check_commentary(ov, pack)
        if probs: die('Commentary override rejected: ' + '; '.join(probs))
        parts = re.split(r'\n\s*(?=UNVERIFIED)', ov, maxsplit=1)
        obs = parts[0].strip(); unv = parts[1].strip() if len(parts) > 1 else ''
        if not unv: die('Commentary override must start the unverified section on a new line beginning "UNVERIFIED"')
    words = len((obs+' '+unv).split())
    probs = check_commentary(obs+'\n'+unv, pack)
    if probs: die('Commentary failed validation: ' + '; '.join(probs))
    es.cell(14, 1, f'CFO commentary ({words} words)').font = fbb
    for r, t in ((15, obs), (16, unv)):
        es.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); c = es.cell(r, 1, t); c.alignment = wrap; c.font = fb; es.row_dimensions[r].height = 62 if r == 15 else 40
    es.cell(17, 1, 'Commentary is generated at build time from verified results and checked (150-word limit, every number traceable to the fact pack); it is static text. KPI tiles above are live formulas.').font = fs
    if bd:
        es['A3'] = ('="Budget: original Sales_Budget is "&TEXT(Budget_Integrity!B5,"0.0")&"x actual revenue and is NOT used. Illustrative variance only against SYNTHETIC v2 (unapproved) - see Budget_vs_Actual_v2 / Budget_Integrity."' if bd['v2'] else '="Budget: original Sales_Budget is "&TEXT(Budget_Integrity!B5,"0.0")&"x actual revenue - see Budget_Integrity."')
        es['A3'].font = Font(name=F, size=10, bold=True, color='C00000')
    es.cell(19, 1, 'Metric availability').font = fbb
    hdr(es, 20, ['Metric','Status','Reason / note'])
    unav = [('Monthly net revenue, COGS, gross profit, GP %','Available','Derived from Sales_Transactions (Actual rows).'),
            ('Region-wise and product-wise performance','Available','Region via Customer_Master; product/category via Product_Master.'),
            ('Gross sales, trade discounts, returns, rebates','UNAVAILABLE','Source revenue is already net (Units x Net_Price); no gross-to-net, returns or rebate fields exist. List_Price vs Net_Price gap is implied only, not a reported discount.'),
            ('GST / intercompany adjustments','UNAVAILABLE','Source README states revenue excludes GST and intercompany sales; no such data supplied.'),
            ('Budget variance (revenue / gross profit)','ILLUSTRATIVE ONLY' if (bd and bd['v2']) else 'UNAVAILABLE (not reported)',('Original budget is not comparable (see Budget_Integrity). Variance exists only against a SYNTHETIC calibrated v2 budget that is unapproved - see Budget_vs_Actual_v2.' if (bd and bd['v2']) else 'No usable budget comparison.')),
            ('Product-level budget','UNAVAILABLE','Budget is by Month x Region x Category only.'),
            ('Region vs product attribution','LIMITED' if ef['confounded'] else 'Available','Each product sells in one region only in this source, so regional and product effects overlap.' if ef['confounded'] else 'Products sell in several regions.'),
            ('Operating profit / EBITDA','OUT OF DAY-1 SCOPE','Opex exists in source but no allocation basis to sales dimensions was provided.')]
    for i, (a, s_, n_) in enumerate(unav):
        r = 21+i; put(es, r, 1, a, al=wrap); put(es, r, 2, s_, font=fbb, al=wrap); put(es, r, 3, n_, al=wrap); es.row_dimensions[r].height = 44
    widths(es, {'A':44,'B':26,'C':96})

    # ---- Validation_Report
    vw = W['Validation_Report']; title(vw, 'Validation Report - what was actually tested', 'Written from the test run on this workbook build. Anything not listed here was NOT tested.')
    hdr(vw, 4, ['#','Test','What was done','Result','Evidence'])
    if validation:
        for i, v in enumerate(validation, 1):
            r = 4+i; put(vw, r, 1, i); put(vw, r, 2, v['test'], al=wrap); put(vw, r, 3, v['method'], al=wrap); put(vw, r, 4, v['result'], font=fbb, al=wrap); put(vw, r, 5, v['evidence'], al=wrap); vw.row_dimensions[r].height = 62
        vw.conditional_formatting.add(f'D5:D{4+len(validation)}', FormulaRule(formula=['LEFT(D5,4)="PASS"'], font=Font(color='006100', bold=True)))
        vw.conditional_formatting.add(f'D5:D{4+len(validation)}', FormulaRule(formula=['LEFT(D5,4)="FAIL"'], font=Font(color='9C0006', bold=True)))
    else:
        put(vw, 5, 2, 'Validation not yet run for this build.')
    widths(vw, {'A':5,'B':34,'C':62,'D':14,'E':62})

    # ---- Data_Dictionary
    dd = W['Data_Dictionary']; title(dd, 'Data Dictionary', 'Source fields used, derived fields and report measures.')
    hdr(dd, 4, ['Sheet','Field','Type','Description','Derivation / rule'])
    rows = [
     ('Sales_Transactions','Invoice_ID','Text','Unique invoice identifier (INV-MM-nnn)','Primary key; uniqueness tested'),
     ('Sales_Transactions','Month','Text YYYY-MM','Reporting month','Links to Calendar.Month'),
     ('Sales_Transactions','Status','Text','Actual or Forecast','Only "Actual" rows enter the MIS; cross-checked to Calendar.Reporting_Status'),
     ('Sales_Transactions','Customer_ID / Product_ID','Text','Foreign keys','To Customer_Master / Product_Master; unmatched keys tested'),
     ('Sales_Transactions','Units','Integer','Units sold','Must be > 0'),
     ('Sales_Transactions','Net_Price_INR / Unit_Cost_INR','Decimal','Net selling price and unit cost per unit, INR','Must be > 0'),
     ('Sales_Transactions','Net_Revenue_INR','Decimal','Net revenue, INR','Source rule: Units x Net_Price_INR (re-performed in controls)'),
     ('Sales_Transactions','COGS_INR','Decimal','Cost of goods sold, INR','Source rule: Units x Unit_Cost_INR (re-performed in controls)'),
     ('Sales_Transactions','Gross_Profit_INR','Decimal','Gross profit, INR','Source rule: Net_Revenue - COGS (re-performed in controls)'),
     ('Product_Master','Product_ID, Product_Name, Category','Text','Product attributes','Lookup for product-wise and category analysis'),
     ('Product_Master','List_Price_INR, Standard_Cost_INR','Decimal','Reference list price / standard cost','Not used in MIS measures'),
     ('Customer_Master','Customer_ID, Customer_Name, Region, Channel','Text','Customer attributes','Region used for region-wise analysis; Channel carried in Calc_Actuals only'),
     ('Calendar','Month, Year, Quarter, Month_Name, Reporting_Status','Text/Integer','Reporting calendar with Actual/Forecast flag','Drives latest actual month and status cross-check'),
     ('Sales_Budget','All fields','Mixed','Budget scenario by Month x Region x Category','Copied only; used for scale observation. Never mixed with actuals'),
     ('Calc_Actuals','Region, Channel, Category, Product_Name','Text','Looked up attributes','INDEX/MATCH to masters; "UNMATCHED" if key missing'),
     ('Calc_Actuals','Year, Month_No, Calendar_Status','Number/Text','Period attributes','From Calendar / RIGHT(Month,2)'),
     ('Calc_Actuals','Units, Net_Revenue_INR, COGS_INR, Gross_Profit_INR','Number','Row-level measures','Direct references to Sales_Transactions'),
     ('Calc_Actuals','Flag_* , Any_Exception','0/1','Row-level exception flags','Unmatched product/customer/month, duplicate invoice, arithmetic, status vs calendar'),
     ('Budget_v2_Synthetic','k','Number','Uniform calibration factor','FY(calibration year) actual revenue / original budget revenue in cells that have sales'),
     ('Budget_v2_Synthetic','In scope (1/0)','0/1','Region x Category cell has at least one actual sale in the source','COUNTIFS on Calc_Actuals'),
     ('Budget_v2_Synthetic','v2 Revenue / COGS / GP','INR','SYNTHETIC calibrated budget (unapproved)','Original budget x k'),
     ('Budget_vs_Actual_v2','Variance','INR / %','Actual - v2 budget, in-scope cells only','Illustrative; calibration-year variance is ~0 by construction'),
     ('Product_Performance','Regions selling (actual)','Integer','Number of regions in which the product has actual sales','COUNTIFS on Calc_Actuals; 1 = region and product fully overlap'),
     ('Monthly_MIS','Net Revenue / COGS','INR','Sum of Actual rows for month','SUMIFS on Calc_Actuals, Status="Actual"'),
     ('Monthly_MIS','Gross Profit','INR','Net Revenue - COGS','Derived; reconciled to source Gross_Profit_INR'),
     ('Monthly_MIS','GP %','%','Gross Profit / Net Revenue','Ratio of sums, never average of ratios'),
     ('Monthly_MIS','MoM / YoY Revenue %','%','Change vs prior month / same month prior year','Shown "n/a" where comparative month is not in the actuals'),
     ('Monthly_MIS','YoY GP % change (pp)','pp','(GP% - prior-year GP%) x 100','Percentage points'),
     ('Region / Product sheets','YTD columns','INR / %','Jan to latest actual month, current vs prior year','Latest actual month read from Calendar (Monthly_MIS!B4)'),
     ('Region / Product sheets','Revenue share','%','Item revenue / total actual revenue','')]
    for i, row in enumerate(rows):
        for j, v in enumerate(row): put(dd, 5+i, 1+j, v, al=wrap)
    widths(dd, {'A':22,'B':36,'C':14,'D':48,'E':64})

    # ---- Automation_Readme
    ar = W['Automation_Readme']; title(ar, 'Automation Readme', 'How this workbook is produced, how to rerun it, and what it does NOT do.')
    sections = [
     ('Source', 'Single input workbook (Nova CFO MIS synthetic dataset). Required sheets/columns are validated before building; a schema failure stops the run before any workbook is produced. The source file is read-only input and is never modified (SHA-256 recorded before and after). Source tables used are copied unchanged into this workbook (grey tabs).'),
     ('Actuals definition', 'Rows with Sales_Transactions.Status = "Actual", cross-checked against Calendar.Reporting_Status (must agree on every row). Forecast and Budget never enter any actual figure.'),
     ('Formulas', 'Net revenue = SUMIFS of Net_Revenue_INR; COGS = SUMIFS of COGS_INR; Gross profit = Net revenue - COGS; GP % = Gross profit / Net revenue (ratio of sums). All report cells are live Excel formulas over Calc_Actuals, which references the raw sheets. Month filters use numeric Year / Month_No (not date-like text) for Excel robustness.'),
     ('Budget handling', 'Original Sales_Budget is carried unchanged and NOT used for variance because it is not comparable (see Budget_Integrity). Budget_v2_Synthetic is a separate, clearly labelled, UNAPPROVED calibrated version used only for illustration in Budget_vs_Actual_v2.'),
     ('One-command run', 'python run_mis.py run --input <dataset.xlsx>  -> validates input, builds the workbook, recalculates it, re-derives every figure independently, applies the publication gate, drafts commentary and writes an edit-detecting hash-chained audit log (not tamper-proof; not externally anchored unless you store its final hash elsewhere) under runs/<run_id>/. Candidate files are NOT published.'),
     ('Publication gate', 'Any failed critical control (workbook controls, independent recompute, error scan, commentary validation, audit-chain integrity, source-unchanged check, and a current passing self-test record bound to the combined hash of run_mis.py + mis_engine.py) sets status BLOCKED; nothing can be approved. A stale, failing or missing self-test record is critical and cannot be overridden by --ack-warnings. Warnings (budget integrity, restatement of previously published months, region/product overlap) require explicit acknowledgement at approval.'),
     ('Human approval', 'python run_mis.py approve --run <run_id> --by "<name>" [--ack-warnings]. Verifies the self-test record, audit chain, candidate manifest, archived input hash and code hash, then copies the candidate into a NEW immutable folder published/<latest-month>/<run_id>/ (never overwrites; prior releases kept; LATEST.json updated only after success). Approval is never automatic. Approver name is typed text, not an authenticated identity. Optional commands: anchor, verify-anchor (external hash record) and verify-release. The script cannot stop files being copied by hand outside the controlled folder: that needs folder permissions or a signed-release process.'),
     ('Dependencies', 'Python 3, pandas, numpy, openpyxl, LibreOffice (headless recalculation). Excel Desktop is NOT used by the pipeline.'),
     ('What is NOT implemented', 'No scheduling, no unattended trigger, no live data connection, no email/distribution, no authenticated approver identity, no cryptographic signing, no external audit anchoring (optional manual step only). The commentary is rule-based text generated from verified facts (an externally written commentary can be supplied with --commentary-file and is checked against the fact pack); no AI model is called by the script. Variance on an original budget is not implemented (only the flagged-budget v2 path). Behaviour in desktop Excel is untested (see Validation_Report).'),
     ('Limitations', 'Commentary and the Budget_Integrity findings table are static text from build time. Revenue excludes GST, returns, rebates and intercompany. The number-grounding check validates magnitudes, not direction or causal wording. Synthetic data only; causes of trends cannot be derived from the dataset.')]
    for i, (a, t) in enumerate(sections):
        r = 4+i; put(ar, r, 1, a, font=fbb, fill=gfill, al=wrap); put(ar, r, 2, t, al=wrap); ar.row_dimensions[r].height = max(34, 15*(len(t)//105+1))
    widths(ar, {'A':24,'B':120})

    for n, col in (('Executive_Summary',NAVY),('Monthly_MIS',NAVY),('Region_Performance',NAVY),('Product_Performance',NAVY),('Control_Checks','548235'),('Validation_Report','548235'),('Data_Dictionary',ORANGE),('Automation_Readme',ORANGE)):
        W[n].sheet_properties.tabColor = col
    wb.save(out); return dict(months=months, m0=m0, m1=m1, mt=mt, PS=PS, RG=RG, PR=PR, nS=nS, nchecks=len(chk), words=words, obs=obs, unv=unv, pack=pack, bd=bd, ef=ef, bva=BVA)

if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='Engine only (no gating/audit). Use run_mis.py for the controlled workflow.')
    ap.add_argument('--input', required=True); ap.add_argument('--output', required=True); ap.add_argument('--validation'); ap.add_argument('--commentary-file')
    a = ap.parse_args()
    v = json.load(open(a.validation)) if a.validation else None
    co = open(a.commentary_file).read() if a.commentary_file else None
    try: info = build(a.input, a.output, v, co)
    except EngineError as e: print('ERROR:', e, file=sys.stderr); sys.exit(2)
    print('Built', a.output, '| actual months:', len(info['months']), '| checks:', info['nchecks'], '| commentary words:', info['words'])
